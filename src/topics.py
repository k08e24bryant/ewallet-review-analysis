"""Phase 6a: BERTopic on complaints across all apps (fit + export for naming).

Documents: unique texts (in_model_pool) that PRIMARY predicts as complaints,
excluding is_short. Every other complaint row (duplicates, short reviews) is
assigned with transform and flagged assigned_by_transform. No gold or dev labels
are read in this phase.

Modes:
    --explore   fit each hdbscan.candidates value; write reports/phase6_candidates.json
                (topic count, outlier share, keywords, outlier similarity curve)
    (default)   final fit with hdbscan.min_cluster_size, optional reduce_outliers,
                naming workbook, assignments for all complaint rows, model, figure,
                report, W&B run
    --smoke     small sample, outputs under models/bertopic_smoke, W&B disabled

Usage:
    uv run python -m src.topics --config configs/topics.yaml --smoke
    uv run python -m src.topics --config configs/topics.yaml --explore
    uv run python -m src.topics --config configs/topics.yaml
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import platform
import sys
import time
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import yaml

from src.gold_report import GRID, SEQ_BLUE, SERIES, SURFACE, TEXT_PRIMARY, TEXT_SECONDARY

logger = logging.getLogger("topics")


def load_config(path: Path) -> dict[str, Any]:
    """Load a YAML config."""
    with path.open(encoding="utf-8") as f:
        return yaml.safe_load(f)


# ---------------------------------------------------------------- data + embeddings

def load_complaints(cfg: dict[str, Any]) -> pd.DataFrame:
    """All complaint rows (PRIMARY) with text and flags; in_fit marks the BERTopic documents."""
    reviews = pd.read_parquet(cfg["input_path"], columns=["reviewId", "app", cfg["text_col"], "is_short", "in_model_pool"])
    pred = pd.read_parquet(cfg["predictions_path"], columns=["reviewId", "is_complaint"])
    df = reviews.merge(pred, on="reviewId", how="inner", validate="one_to_one")
    if len(df) != len(reviews):
        raise ValueError("predictions do not cover every review; rerun src.predict")
    df = df[df["is_complaint"]].drop(columns="is_complaint").reset_index(drop=True)
    df["in_fit"] = df["in_model_pool"] & ~df["is_short"]
    return df


def embed(texts: list[str], cfg: dict[str, Any], cache_dir: Path) -> np.ndarray:
    """Sentence embeddings for unique texts, cached to disk by model name + text hash."""
    e = cfg["embedding"]
    digest = hashlib.sha256("\x1f".join(texts).encode("utf-8")).hexdigest()[:16]
    path = cache_dir / f"{e['model'].split('/')[-1]}_{len(texts)}_{digest}.npy"
    if path.exists():
        logger.info("embeddings: cache hit %s", path)
        return np.load(path)
    from sentence_transformers import SentenceTransformer

    import torch

    device = "cuda" if torch.cuda.is_available() else "cpu"
    model = SentenceTransformer(e["model"], device=device)
    t0 = time.time()
    emb = model.encode(texts, batch_size=e["batch_size"], show_progress_bar=False, convert_to_numpy=True)
    logger.info("embedded %d texts on %s in %.0fs", len(texts), device, time.time() - t0)
    cache_dir.mkdir(parents=True, exist_ok=True)
    np.save(path, emb.astype(np.float32))
    return emb.astype(np.float32)


# ---------------------------------------------------------------- model

def stopwords(cfg: dict[str, Any]) -> list[str]:
    """Indonesian stopwords (stopwordsiso) plus the configured slang fillers and app words."""
    import stopwordsiso

    v = cfg["vectorizer"]
    return sorted(set(stopwordsiso.stopwords(v["stopwords"])) | set(v["extra_stopwords"]))


def build_model(cfg: dict[str, Any], min_cluster_size: int):
    """BERTopic with seeded UMAP, HDBSCAN (prediction_data for transform), and the configured vectorizer."""
    from bertopic import BERTopic
    from hdbscan import HDBSCAN
    from sklearn.feature_extraction.text import CountVectorizer
    from umap import UMAP

    u, h, v = cfg["umap"], cfg["hdbscan"], cfg["vectorizer"]
    return BERTopic(
        embedding_model=cfg["embedding"]["model"],
        umap_model=UMAP(n_neighbors=u["n_neighbors"], n_components=u["n_components"], min_dist=u["min_dist"],
                        metric=u["metric"], random_state=u["random_state"], low_memory=False),
        hdbscan_model=HDBSCAN(min_cluster_size=min_cluster_size, metric=h["metric"],
                              cluster_selection_method=h["cluster_selection_method"], prediction_data=True),
        vectorizer_model=CountVectorizer(ngram_range=tuple(v["ngram_range"]), min_df=v["min_df"], stop_words=stopwords(cfg)),
        top_n_words=cfg["top_n_words"],
        calculate_probabilities=False,
        verbose=False,
    )


def keywords(model, topic: int, n: int) -> list[str]:
    """Top-n c-TF-IDF words of a topic."""
    return [w for w, _ in (model.get_topic(topic) or [])][:n]


def outlier_similarity(model, emb: np.ndarray, topics: np.ndarray) -> np.ndarray:
    """Max cosine similarity of each outlier doc to any topic embedding (what reduce_outliers 'embeddings' uses)."""
    from sklearn.metrics.pairwise import cosine_similarity

    te = np.asarray(model.topic_embeddings_)[model._outliers:]
    return cosine_similarity(emb[topics == -1], te).max(axis=1)


def fit_one(cfg: dict[str, Any], docs: list[str], emb: np.ndarray, mcs: int) -> tuple[Any, np.ndarray, dict[str, Any]]:
    """Fit BERTopic once; return model, topics, and a summary (count, outlier share, sizes, outlier curve)."""
    t0 = time.time()
    model = build_model(cfg, mcs)
    topics, _ = model.fit_transform(docs, embeddings=emb)
    topics = np.asarray(topics)
    sim = outlier_similarity(model, emb, topics)
    n_out = int((topics == -1).sum())
    sizes = pd.Series(topics[topics != -1]).value_counts()
    summary = {
        "min_cluster_size": mcs, "n_topics": int(sizes.size), "outlier_share": n_out / len(topics),
        "largest_topic_share": float(sizes.max() / len(topics)) if sizes.size else 0.0,
        "smallest_topic_size": int(sizes.min()) if sizes.size else 0,
        "outlier_share_after_reduction": {
            str(t): (n_out - int((sim >= t).sum())) / len(topics) for t in cfg["outliers"]["thresholds_to_report"]},
        "fit_seconds": round(time.time() - t0, 1),
    }
    return model, topics, summary


# ---------------------------------------------------------------- outputs

def topic_table(model, docs: list[str], emb: np.ndarray, topics: np.ndarray, apps: np.ndarray, cfg: dict[str, Any]) -> pd.DataFrame:
    """One row per topic (outliers first): size, share of each app's fitted complaints, keywords, representatives."""
    n_rep = cfg["export"]["n_representative"]
    app_totals = pd.Series(apps).value_counts()
    norm = emb / np.linalg.norm(emb, axis=1, keepdims=True)
    rows = []
    for t in sorted(set(topics)):
        idx = np.flatnonzero(topics == t)
        centroid = norm[idx].mean(0)
        nearest = idx[np.argsort(-(norm[idx] @ centroid))[:n_rep]]
        row = {"topic_id": int(t), "size": int(len(idx))}
        for a in cfg["apps"]:
            row[f"share_{a}"] = float((apps[idx] == a).sum() / app_totals.get(a, np.nan))
        row["keywords"] = ", ".join(keywords(model, t, cfg["top_n_words"])) if t != -1 else ""
        for i, j in enumerate(nearest, 1):
            row[f"rep_{i}"] = docs[j]
        rows.append(row)
    return pd.DataFrame(rows)


def write_workbook(table: pd.DataFrame, path: Path, cfg: dict[str, Any]) -> None:
    """topics_to_name.xlsx: one row per topic + empty topic_name / merge_into / notes; refuses to overwrite."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    if path.exists():
        raise FileExistsError(f"{path} exists; refusing to overwrite (it may contain your topic names)")
    path.parent.mkdir(parents=True, exist_ok=True)
    df = table.copy()
    df["topic_name"], df["merge_into"], df["notes"] = None, None, None
    df.loc[df["topic_id"] == -1, "topic_name"] = "(outliers: not named)"
    wb = Workbook()
    ws = wb.active
    ws.title = "topics"
    ws.append(list(df.columns))
    for r in df.itertuples(index=False):
        ws.append(list(r))
    widths = {"topic_id": 9, "size": 8, "keywords": 40, "topic_name": 28, "merge_into": 12, "notes": 30}
    for i, col in enumerate(df.columns, 1):
        letter = ws.cell(1, i).column_letter
        ws.column_dimensions[letter].width = widths.get(col, 11 if col.startswith("share_") else 45)
        ws.cell(1, i).font = Font(bold=True)
        for row in range(2, len(df) + 2):
            c = ws.cell(row, i)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if col.startswith("share_"):
                c.number_format = "0.0%"
    user_fill = PatternFill("solid", fgColor="FFF6D5")
    gray = PatternFill("solid", fgColor="E4E3DF")
    for col in ("topic_name", "merge_into", "notes"):
        ci = list(df.columns).index(col) + 1
        ws.cell(1, ci).fill = user_fill
    for row in range(2, len(df) + 2):
        if ws.cell(row, 1).value == -1:
            for ci in range(1, len(df.columns) + 1):
                ws.cell(row, ci).fill = gray
    ids = [str(t) for t in df["topic_id"] if t != -1]
    mi = ws.cell(1, list(df.columns).index("merge_into") + 1).column_letter
    dv = DataValidation(type="list", formula1=f'"{",".join(ids)}"' if len(",".join(ids)) < 250 else None,
                        allow_blank=True, showErrorMessage=True, error="merge_into must be an existing topic_id")
    if dv.formula1:
        ws.add_data_validation(dv)
        dv.add(f"{mi}2:{mi}{len(df) + 1}")
    ws.freeze_panes = "C2"

    info = wb.create_sheet("README")
    for line in [
        "One row per topic from BERTopic on complaint reviews (unique texts, not short).",
        "size: fitted documents in the topic.",
        "share_<app>: the topic's share of that app's fitted complaint documents (columns of one app sum to 1 with outliers).",
        "keywords: top c-TF-IDF words (Indonesian stopwords, slang fillers and app names removed).",
        "rep_1..rep_5: the reviews closest to the topic's embedding centroid.",
        "Fill in: topic_name (short label), merge_into (topic_id to merge this topic into; leave blank to keep), notes.",
        "Topic -1 = outliers; listed for reference, not named.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 120
    wb.save(path)


def plot_overview(table: pd.DataFrame, path: Path, cfg: dict[str, Any], subtitle: str) -> None:
    """Left: topic size; right: each topic's share of each app's complaints (one sequential hue)."""
    import matplotlib.pyplot as plt
    from matplotlib.colors import LinearSegmentedColormap

    t = table[table["topic_id"] != -1].sort_values("topic_id").reset_index(drop=True)
    labels = [f"T{r.topic_id}  " + " · ".join(r.keywords.split(", ")[:3]) for r in t.itertuples()]
    shares = t[[f"share_{a}" for a in cfg["apps"]]].to_numpy()
    h = 0.32 * len(t) + 1.8
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(13, h), facecolor=SURFACE, sharey=True,
                                   gridspec_kw={"width_ratios": [1.1, 1.0]})
    y = np.arange(len(t))
    ax1.set_facecolor(SURFACE)
    ax1.barh(y, t["size"], color=SERIES[0], height=0.72, edgecolor=SURFACE, linewidth=2, zorder=2)
    for yi, v in zip(y, t["size"]):
        ax1.text(v, yi, f" {v:,}", va="center", color=TEXT_SECONDARY, fontsize=8)
    ax1.set_yticks(y, labels, fontsize=8.5)
    ax1.set_ylim(len(t) - 0.5, -0.5)
    ax1.grid(axis="x", color=GRID, linewidth=0.8, zorder=0)
    for s in ("top", "right", "left"):
        ax1.spines[s].set_visible(False)
    ax1.spines["bottom"].set_color(GRID)
    ax1.tick_params(colors=TEXT_SECONDARY, length=0)
    ax1.set_xlabel("Documents (unique complaint texts)", color=TEXT_SECONDARY)
    ax1.set_xlim(0, t["size"].max() * 1.18)

    cmap = LinearSegmentedColormap.from_list("seq_blue", ["#f3f7fd"] + SEQ_BLUE)
    vmax = max(shares.max(), 1e-9)
    ax2.imshow(shares, cmap=cmap, vmin=0, vmax=vmax, aspect="auto")
    for i in range(shares.shape[0]):
        for j in range(shares.shape[1]):
            v = shares[i, j]
            ax2.text(j, i, f"{v:.0%}" if v >= 0.005 else "", ha="center", va="center", fontsize=7.5,
                     color="white" if v / vmax >= 0.55 else TEXT_PRIMARY)
    ax2.set_xticks(range(len(cfg["apps"])), cfg["apps"])
    ax2.xaxis.tick_top()
    ax2.tick_params(colors=TEXT_SECONDARY, length=0)
    for s in ax2.spines.values():
        s.set_visible(False)
    ax2.set_xticks(np.arange(-0.5, len(cfg["apps"])), minor=True)
    ax2.set_yticks(np.arange(-0.5, len(t)), minor=True)
    ax2.grid(which="minor", color=SURFACE, linewidth=2)
    ax2.tick_params(which="minor", length=0)
    ax2.set_title("Share of each app's complaints", color=TEXT_PRIMARY, fontsize=10, loc="left", pad=22)
    fig.text(0.01, 0.995, "Phase 6a: complaint topics across apps (unnamed; labels = top keywords)",
             color=TEXT_PRIMARY, fontsize=12, va="top")
    fig.text(0.01, 0.995 - 0.3 / h, subtitle, color=TEXT_SECONDARY, fontsize=9, va="top")
    fig.tight_layout(rect=(0, 0, 1, 1 - 0.6 / h))
    fig.savefig(path, dpi=150, facecolor=SURFACE)
    plt.close(fig)


# ---------------------------------------------------------------- main

def explore(cfg: dict[str, Any], docs: list[str], emb: np.ndarray, out: Path) -> None:
    """Fit each candidate min_cluster_size and write the comparison."""
    results = []
    for mcs in cfg["hdbscan"]["candidates"]:
        logger.info("fitting candidate min_cluster_size=%d", mcs)
        model, topics, s = fit_one(cfg, docs, emb, mcs)
        s["topics"] = {int(t): {"size": int(n), "keywords": keywords(model, t, 8)}
                       for t, n in pd.Series(topics[topics != -1]).value_counts().sort_index().items()}
        results.append(s)
        print(f"\n== min_cluster_size={mcs}: {s['n_topics']} topics, outliers {s['outlier_share']:.1%}, "
              f"largest {s['largest_topic_share']:.1%}, smallest {s['smallest_topic_size']} ({s['fit_seconds']}s)")
        print("   outliers left after reduce_outliers at threshold:",
              {k: f"{v:.1%}" for k, v in s["outlier_share_after_reduction"].items()})
        for t, d in s["topics"].items():
            print(f"   T{t:<3} {d['size']:>6}  {', '.join(d['keywords'])}")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps({"seed": cfg["seed"], "n_docs": len(docs), "candidates": results}, indent=2,
                              ensure_ascii=False), encoding="utf-8")
    logger.info("wrote %s", out)


def main() -> None:  # noqa: PLR0915 - linear pipeline
    """CLI entry point."""
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--config", type=Path, default=Path("configs/topics.yaml"))
    parser.add_argument("--explore", action="store_true", help="fit each candidate min_cluster_size and stop")
    parser.add_argument("--smoke", action="store_true", help="small sample, outputs under models/bertopic_smoke")
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")
    for noisy in ("httpx", "huggingface_hub", "BERTopic", "numba"):
        logging.getLogger(noisy).setLevel(logging.WARNING)

    cfg = load_config(args.config)
    seed = cfg["seed"]
    np.random.seed(seed)
    text = cfg["text_col"]
    df = load_complaints(cfg)
    fit_df = df[df["in_fit"]].reset_index(drop=True)
    other = df[~df["in_fit"]].reset_index(drop=True)
    paths = {k: Path(cfg[k]) for k in ("candidates_path", "report_path", "assignments_path", "model_dir")}
    paths["xlsx"] = Path(cfg["export"]["xlsx_path"])
    paths["figure"] = Path(cfg["figures_dir"]) / cfg["figure"]
    mcs = cfg["hdbscan"]["min_cluster_size"]
    threshold = cfg["outliers"]["threshold"]
    if args.smoke:
        fit_df = fit_df.sample(n=cfg["smoke"]["fit_docs"], random_state=seed).reset_index(drop=True)
        other = other.sample(n=cfg["smoke"]["transform_rows"], random_state=seed).reset_index(drop=True)
        out = Path(cfg["smoke"]["out_dir"])
        paths = {"candidates_path": out / "candidates.json", "report_path": out / "report.json",
                 "assignments_path": out / "assignments.parquet", "model_dir": out / "model",
                 "xlsx": out / "topics_to_name.xlsx", "figure": out / cfg["figure"]}
        cfg["hdbscan"]["candidates"] = [cfg["smoke"]["min_cluster_size"]]
        mcs = cfg["smoke"]["min_cluster_size"]
        threshold = threshold if threshold is not None else 0.4
        for p in (paths["xlsx"],):
            p.unlink(missing_ok=True)  # smoke outputs only
    logger.info("seed=%d smoke=%s complaint rows=%d fit docs=%d transform rows=%d", seed, args.smoke, len(df),
                len(fit_df), len(other))

    # Embeddings for every unique text that needs a topic (fit docs + transform rows)
    uniq = pd.Index(pd.concat([fit_df[text], other[text]]).drop_duplicates())
    emb_all = embed(uniq.tolist(), cfg, Path(cfg["embedding"]["cache_dir"]) if not args.smoke else Path(cfg["smoke"]["out_dir"]) / "emb")
    docs = fit_df[text].tolist()
    emb = emb_all[uniq.get_indexer(docs)]

    if args.explore:
        explore(cfg, docs, emb, paths["candidates_path"])
        return
    if mcs is None:
        raise ValueError("set hdbscan.min_cluster_size (and outliers.threshold) after --explore")

    # ---- final fit
    if args.smoke:
        explore(cfg, docs, emb, paths["candidates_path"])
    logger.info("final fit: min_cluster_size=%d on %d docs", mcs, len(docs))
    model, topics_raw, fit_summary = fit_one(cfg, docs, emb, mcs)
    out_cfg = cfg["outliers"]

    # ---- transform every other complaint row (duplicates, short), before any topic update
    t0 = time.time()
    o_texts = other[text].drop_duplicates().tolist()
    o_emb = emb_all[uniq.get_indexer(o_texts)]
    o_raw = np.asarray(model.transform(o_texts, embeddings=o_emb)[0]) if o_texts else np.array([], dtype=int)
    logger.info("transform: %d unique texts in %.0fs", len(o_texts), time.time() - t0)

    # Transform consistency on fit docs (raw topics, before outlier reduction)
    rng = np.random.default_rng(seed)
    chk = rng.choice(len(docs), size=min(cfg["transform_check_n"], len(docs)), replace=False)
    chk_raw = np.asarray(model.transform([docs[i] for i in chk], embeddings=emb[chk])[0])
    transform_check = {"n": int(len(chk)), "agreement_with_fit_topic": float(np.mean(chk_raw == topics_raw[chk])),
                       "agreement_excluding_fit_outliers": float(np.mean((chk_raw == topics_raw[chk])[topics_raw[chk] != -1]))}

    # ---- outlier reduction: fit docs and transform rows against the same (centroid) topic embeddings,
    # then refresh keywords; update_topics replaces topic embeddings, so it must come last
    topics, o_top = topics_raw.copy(), o_raw
    before = float((topics == -1).mean())
    reduction = {"applied": False, "outlier_share_before": before, "outlier_share_after": before}
    if before > out_cfg["max_share"]:
        if threshold is None:
            raise ValueError(f"outlier share {before:.1%} > {out_cfg['max_share']:.0%}: set outliers.threshold")
        topics = np.asarray(model.reduce_outliers(docs, topics_raw.tolist(), strategy=out_cfg["strategy"],
                                                  embeddings=emb, threshold=threshold))
        if len(o_texts):
            o_top = np.asarray(model.reduce_outliers(o_texts, o_raw.tolist(), strategy=out_cfg["strategy"],
                                                     embeddings=o_emb, threshold=threshold))
        model.update_topics(docs, topics=topics.tolist(), vectorizer_model=model.vectorizer_model,
                            top_n_words=cfg["top_n_words"])
        reduction = {"applied": True, "strategy": out_cfg["strategy"], "threshold": threshold,
                     "outlier_share_before": before, "outlier_share_after": float((topics == -1).mean()),
                     "reassigned_docs": int(((topics_raw == -1) & (topics != -1)).sum())}
        logger.info("reduce_outliers: %.1f%% -> %.1f%%", 100 * before, 100 * reduction["outlier_share_after"])

    # ---- assignments for all complaint rows
    fit_topic = pd.Series(topics, index=fit_df["reviewId"])
    fit_raw = pd.Series(topics_raw, index=fit_df["reviewId"])
    o_map, o_map_raw = dict(zip(o_texts, o_top)), dict(zip(o_texts, o_raw))
    assign = pd.concat([
        fit_df.assign(topic=fit_df["reviewId"].map(fit_topic), topic_before_outlier_reduction=fit_df["reviewId"].map(fit_raw),
                      assigned_by_transform=False),
        other.assign(topic=other[text].map(o_map), topic_before_outlier_reduction=other[text].map(o_map_raw),
                     assigned_by_transform=True),
    ], ignore_index=True)
    assign["topic"] = assign["topic"].astype(int)
    assign["topic_before_outlier_reduction"] = assign["topic_before_outlier_reduction"].astype(int)
    assign = assign[["reviewId", "app", "topic", "topic_before_outlier_reduction", "assigned_by_transform",
                     "in_fit", "in_model_pool", "is_short"]]
    assert len(assign) == len(fit_df) + len(other) and assign["reviewId"].is_unique
    # Duplicate rows whose text is also a fit doc: does transform reproduce the fitted topic?
    text_topic = dict(zip(docs, topics))
    dup = other[other[text].isin(text_topic)]
    dup_agree = float(np.mean(dup[text].map(o_map).to_numpy() == dup[text].map(text_topic).to_numpy())) if len(dup) else None
    assign_stats = {
        "rows": int(len(assign)), "fit": int((~assign["assigned_by_transform"]).sum()),
        "transform": int(assign["assigned_by_transform"].sum()),
        "transform_by_reason": {"short": int((other["is_short"]).sum()),
                                "duplicate_text_not_short": int((~other["is_short"] & ~other["in_model_pool"]).sum())},
        "outlier_share_transform_rows_before": float(np.mean(other[text].map(o_map_raw) == -1)) if len(other) else None,
        "outlier_share_transform_rows_after": float(np.mean(other[text].map(o_map) == -1)) if len(other) else None,
        "outlier_share_all_rows": float((assign["topic"] == -1).mean()),
        "duplicate_rows_with_fit_text": int(len(dup)), "duplicate_transform_matches_fit_topic": dup_agree,
    }

    # ---- table, workbook, figure, model
    table = topic_table(model, docs, emb, topics, fit_df["app"].to_numpy(), cfg)
    write_workbook(table, paths["xlsx"], cfg)
    paths["assignments_path"].parent.mkdir(parents=True, exist_ok=True)
    assign.to_parquet(paths["assignments_path"], index=False)
    paths["figure"].parent.mkdir(parents=True, exist_ok=True)
    n_topics = int((table["topic_id"] != -1).sum())
    plot_overview(table, paths["figure"], cfg,
                  f"{len(docs):,} unique complaint texts (not short), {n_topics} topics, min_cluster_size={mcs}; "
                  f"outliers {reduction['outlier_share_after']:.0%} (before reduction {reduction['outlier_share_before']:.0%}).")
    paths["model_dir"].mkdir(parents=True, exist_ok=True)
    model.save(str(paths["model_dir"] / "bertopic.pkl"), serialization="pickle", save_embedding_model=False)

    import bertopic
    import hdbscan
    import umap

    pred_report = json.loads(Path(cfg["predictions_report"]).read_text(encoding="utf-8")) if not args.smoke else None
    dana_cicil = int(sum("cicil" in d.lower() for d in docs))
    report = {
        "seed": seed, "smoke": args.smoke, "labels_used": "none (PRIMARY predictions only; no gold or dev)",
        "predictions": pred_report,
        "documents": {"definition": "in_model_pool & PRIMARY complaint & not is_short; text_clean",
                      "n": len(docs), "per_app": fit_df["app"].value_counts().to_dict(),
                      "complaint_rows_total": int(len(df))},
        "embedding": {**cfg["embedding"], "dim": int(emb.shape[1])},
        "umap": cfg["umap"], "hdbscan": {**cfg["hdbscan"], "min_cluster_size": mcs},
        "vectorizer": {**cfg["vectorizer"], "n_stopwords": len(stopwords(cfg))},
        "candidates_report": str(paths["candidates_path"]),
        "fit": fit_summary,
        "outlier_reduction": reduction,
        "transform_check_on_fit_docs": transform_check,
        "assignments": assign_stats,
        "docs_mentioning_cicil": dana_cicil,
        "topics": table.drop(columns=[c for c in table.columns if c.startswith("rep_")]).to_dict("records"),
        "outputs": {"workbook": str(paths["xlsx"]), "assignments": str(paths["assignments_path"]),
                    "model": str(paths["model_dir"] / "bertopic.pkl"), "figure": str(paths["figure"])},
        "versions": {"python": platform.python_version(), "bertopic": bertopic.__version__,
                     "umap": umap.__version__, "hdbscan": getattr(hdbscan, "__version__", "?")},
    }
    paths["report_path"].parent.mkdir(parents=True, exist_ok=True)
    paths["report_path"].write_text(json.dumps(report, indent=2, ensure_ascii=False, default=str), encoding="utf-8")

    report["wandb_url"] = log_wandb(report, table, paths["figure"], cfg, args.smoke)
    paths["report_path"].write_text(json.dumps(report, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    logger.info("wrote %s, %s, %s, %s", paths["xlsx"], paths["assignments_path"], paths["report_path"], paths["figure"])

    print(f"\n== final: {n_topics} topics on {len(docs):,} docs; outliers {before:.1%} -> {reduction['outlier_share_after']:.1%}")
    print("transform check on fit docs:", transform_check)
    print("assignments:", assign_stats)
    for r in table.itertuples():
        print(f"   T{r.topic_id:<3} {r.size:>6}  " + "  ".join(f"{a[:4]} {getattr(r, f'share_{a}'):.0%}" for a in cfg["apps"])
              + f"  | {r.keywords}")


def log_wandb(report: dict[str, Any], table: pd.DataFrame, figure: Path, cfg: dict[str, Any], smoke: bool) -> str | None:
    """One W&B run: config, topic table, overview figure, headline numbers."""
    import wandb

    from src.train_indobert import wandb_logged_in

    icfg = load_config(Path("configs/indobert.yaml"))
    if not smoke and not wandb_logged_in():
        logger.warning("W&B not logged in; skipping")
        return None
    try:
        run = wandb.init(project=icfg["wandb"]["project"], entity=icfg["wandb"]["entity"], name=cfg["wandb"]["name"],
                         job_type=cfg["wandb"]["job_type"], tags=["phase6", "bertopic"],
                         config={"seed": cfg["seed"], "embedding": cfg["embedding"]["model"], "umap": cfg["umap"],
                                 "hdbscan": report["hdbscan"], "vectorizer": report["vectorizer"],
                                 "outliers": report["outlier_reduction"]},
                         mode="disabled" if smoke else None)
        t = table.drop(columns=[c for c in table.columns if c.startswith("rep_")])
        run.log({"topics": wandb.Table(dataframe=t), "topic_overview": wandb.Image(str(figure))})
        run.summary.update({"n_docs": report["documents"]["n"], "n_topics": int((t["topic_id"] != -1).sum()),
                            "outlier_share_before": report["outlier_reduction"]["outlier_share_before"],
                            "outlier_share_after": report["outlier_reduction"]["outlier_share_after"],
                            "transform_agreement_fit_docs": report["transform_check_on_fit_docs"]["agreement_with_fit_topic"]})
        url = None if smoke else run.url
        run.finish()
        return url
    except Exception as exc:  # noqa: BLE001 - outputs are saved; W&B is a mirror
        logger.warning("W&B logging failed: %s", exc)
        return f"failed: {exc}"


if __name__ == "__main__":
    main()
