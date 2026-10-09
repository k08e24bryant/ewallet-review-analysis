"""Phase 6c-1: improved topic model (v2) after the low 6b fit rate.

The 60 Phase 6b fit-check rows are diagnostic only; v2 is judged on a fresh sample
in 6c-2 (adopt v2 if its fresh "yes" rate is above 37%). No gold or dev labels.

Modes:
    --compare   embed the 60,014 fit documents with each configured model (cached) and fit
                BERTopic with identical settings; compare NPMI coherence, outlier share,
                and writing-style topics -> reports/phase6c_embedding_comparison.json
    --fit NAME  refit with the chosen embedding, calibrate outlier reduction to the 6a
                outlier level, save the model, draw 20 random documents per topic to read
                (their reviewIds are saved so 6c-2 can exclude them), print them
    --export    build data/topics/topics_to_name_<version>.xlsx from my proposals
                (data/topics/topic_proposals_<version>.yaml); refuses to overwrite
    --version   v2 (default) or v3 (brand names neutralized in the embedding input only)

Usage:
    uv run python -m src.topics_v2 --compare
    uv run python -m src.topics_v2 --fit indo_sbert
    uv run python -m src.topics_v2 --export
    uv run python -m src.topics_v2 --version v3 --fit indo_sbert
    uv run python -m src.topics_v2 --version v3 --export
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import sys
import time
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import yaml

from src import topics as tp

logger = logging.getLogger("topics_v2")


def load_config(path: Path) -> dict[str, Any]:
    """Load a YAML config."""
    with path.open(encoding="utf-8") as f:
        return yaml.safe_load(f)


# ---------------------------------------------------------------- embeddings + model

def embed(texts: list[str], spec: dict[str, Any], cfg: dict[str, Any]) -> np.ndarray:
    """Embeddings for texts with the model's prefix and a fixed max length; cached to disk."""
    key = f"{spec['model']}|{spec['prefix']}|{cfg['max_seq_length']}|" + "\x1f".join(texts)
    digest = hashlib.sha256(key.encode("utf-8")).hexdigest()[:16]
    path = Path(cfg["cache_dir"]) / f"{spec['name']}_{len(texts)}_{digest}.npy"
    if path.exists():
        logger.info("embeddings %s: cache hit %s", spec["name"], path)
        return np.load(path)
    import torch
    from sentence_transformers import SentenceTransformer

    device = "cuda" if torch.cuda.is_available() else "cpu"
    model = SentenceTransformer(spec["model"], device=device)
    model.max_seq_length = cfg["max_seq_length"]
    t0 = time.time()
    emb = model.encode([spec["prefix"] + t for t in texts], batch_size=cfg["batch_size"],
                       show_progress_bar=False, convert_to_numpy=True).astype(np.float32)
    logger.info("embedded %d texts with %s on %s in %.0fs (dim %d)", len(texts), spec["name"], device,
                time.time() - t0, emb.shape[1])
    path.parent.mkdir(parents=True, exist_ok=True)
    np.save(path, emb)
    del model
    if device == "cuda":
        torch.cuda.empty_cache()
    return emb


def stopwords_v2(tcfg: dict[str, Any], fcfg: dict[str, Any]) -> list[str]:
    """6a stopwords + 6b extra stopwords; keep_words must not be in the list."""
    stop = sorted(set(tp.stopwords(tcfg)) | set(fcfg["extra_stopwords"]))
    clash = set(stop) & set(fcfg["keep_words"])
    if clash:
        raise ValueError(f"keep_words are in the stopword list: {clash}")
    return stop


def build(tcfg: dict[str, Any], fcfg: dict[str, Any], cfg: dict[str, Any], spec: dict[str, Any], mcs: int):
    """BERTopic as in 6a (seeded UMAP, HDBSCAN) with the v2 vectorizer and the given embedding model."""
    from sklearn.feature_extraction.text import CountVectorizer

    m = tp.build_model({**tcfg, "top_n_words": cfg["top_n_words"]}, mcs)
    v = tcfg["vectorizer"]
    m.vectorizer_model = CountVectorizer(ngram_range=tuple(v["ngram_range"]), min_df=v["min_df"],
                                         stop_words=stopwords_v2(tcfg, fcfg))
    m.embedding_model = spec["model"]
    return m


# ---------------------------------------------------------------- metrics

def npmi_per_topic(model, docs: list[str], stop: list[str], ngram: tuple[int, int]) -> dict[int, float]:
    """Mean pairwise NPMI of each topic's top-10 keywords, with documents as co-occurrence windows.

    A pair that never co-occurs scores -1. Keywords missing from the documents are skipped.
    """
    from sklearn.feature_extraction.text import CountVectorizer

    topics = sorted(t for t in set(model.topics_) if t != -1)
    kws = {t: tp.keywords(model, t, 10) for t in topics}
    vocab = sorted({w for ws in kws.values() for w in ws})
    vec = CountVectorizer(ngram_range=ngram, stop_words=stop, vocabulary=vocab, binary=True)
    X = vec.fit_transform(docs).tocsc().astype(np.float64)
    n = X.shape[0]
    col = {w: i for i, w in enumerate(vec.get_feature_names_out())}
    out = {}
    for t, ws in kws.items():
        idx = [col[w] for w in ws if w in col and X[:, col[w]].nnz > 0]
        if len(idx) < 2:
            out[t] = float("nan")
            continue
        sub = X[:, idx]
        co = (sub.T @ sub).toarray() / n
        p = np.diag(co)
        vals = []
        for i in range(len(idx)):
            for j in range(i + 1, len(idx)):
                pij = co[i, j]
                vals.append(-1.0 if pij == 0 else np.log(pij / (p[i] * p[j])) / -np.log(pij))
        out[t] = float(np.mean(vals))
    return out


def style_topics(model, lexicon: set[str], threshold: int) -> list[int]:
    """Topics whose top-10 keywords contain at least `threshold` style words (unigrams or parts of bigrams)."""
    flagged = []
    for t in sorted(t for t in set(model.topics_) if t != -1):
        hits = sum(any(part in lexicon for part in w.split()) for w in tp.keywords(model, t, 10))
        if hits >= threshold:
            flagged.append(int(t))
    return flagged


def summarize(model, docs: list[str], cfg: dict[str, Any], stop: list[str], tcfg: dict[str, Any]) -> dict[str, Any]:
    """Topic count, outlier share, NPMI (mean and size-weighted), style topics, largest topic share."""
    topics = np.asarray(model.topics_)
    npmi = npmi_per_topic(model, docs, stop, tuple(tcfg["vectorizer"]["ngram_range"]))
    sizes = pd.Series(topics[topics != -1]).value_counts()
    s = pd.Series(npmi).dropna()
    style = style_topics(model, set(cfg["style_lexicon"]), cfg["style_threshold"])
    return {
        "n_topics": int(sizes.size), "outlier_share": float((topics == -1).mean()),
        "largest_topic_share": float(sizes.max() / len(topics)),
        "npmi_mean": float(s.mean()), "npmi_size_weighted": float((s * sizes.reindex(s.index)).sum() / sizes.reindex(s.index).sum()),
        "style_topics": style, "n_style_topics": len(style),
        "style_docs_share": float(sizes.reindex(style).sum() / len(topics)) if style else 0.0,
        "topics": {int(t): {"size": int(sizes[t]), "npmi": round(npmi[t], 3), "keywords": tp.keywords(model, t, 10)}
                   for t in sizes.sort_index().index},
    }


# ---------------------------------------------------------------- brand neutralization (v3)

def neutralizer(ncfg: dict[str, Any]):
    """Return (fn, patterns): fn(text) -> (embedding input, {pattern_name: replacements})."""
    import re

    clit = "|".join(ncfg["clitics"])
    rules = [(f"paylater:{p}", re.compile(rf"\b{p}(?=(?:{clit})?\b)", re.IGNORECASE), "paylater")
             for p in ncfg["paylater"]]
    rules += [(f"brand:{p}", re.compile(rf"\b{p}(?=(?:{clit})?\b)", re.IGNORECASE), ncfg["token"])
              for p in ncfg["brands"]]

    def fn(text: str) -> tuple[str, dict[str, int]]:
        counts = {}
        for name, rx, repl in rules:
            text, n = rx.subn(repl, text)
            if n:
                counts[name] = n
        return text, counts

    return fn


def neutralize_all(docs: list[str], ncfg: dict[str, Any]) -> tuple[list[str], dict[str, Any]]:
    """Embedding inputs with brands/paylater normalized, plus replacement and residual-typo counts."""
    import re

    fn = neutralizer(ncfg)
    out, totals, changed = [], {}, 0
    for d in docs:
        t, cnt = fn(d)
        out.append(t)
        changed += bool(cnt)
        for k, v in cnt.items():
            totals[k] = totals.get(k, 0) + v
    resid = {w: int(sum(bool(re.search(rf"\b{w}\b", d, re.IGNORECASE)) for d in out)) for w in ncfg["residual_typos"]}
    return out, {"docs_changed": changed, "docs_changed_share": changed / len(docs), "replacements": totals,
                 "residual_typo_docs": resid,
                 "examples": [[a, b] for a, b in zip(docs, out) if a != b][:8]}


def app_profile(apps: np.ndarray, topics: np.ndarray, min_share: float = 0.8) -> dict[str, Any]:
    """Per topic: dominant app, its share of the topic, and the topic's share of that app's documents.

    Topics where one app has >= min_share of the documents are listed as app-dominated.
    """
    base = pd.Series(apps).value_counts()
    rows = {}
    for t in sorted(set(topics) - {-1}):
        a = pd.Series(apps[topics == t]).value_counts()
        dom = a.index[0]
        rows[int(t)] = {"size": int(a.sum()), "dominant_app": dom, "dominant_share": round(float(a.iloc[0] / a.sum()), 3),
                        "share_of_app_docs": round(float(a.iloc[0] / base[dom]), 3)}
    dominated = {t: r for t, r in rows.items() if r["dominant_share"] >= min_share}
    return {"min_share": min_share, "topics": rows, "app_dominated": dominated,
            "app_dominated_docs_share_by_app": {
                a: round(float(sum(r["share_of_app_docs"] for r in dominated.values() if r["dominant_app"] == a)), 3)
                for a in base.index}}


def unspecified_shares(apps: np.ndarray, topics: np.ndarray, proposals: dict[int, dict[str, Any]],
                       bucket: str = "Unspecified complaint") -> dict[str, Any]:
    """Share of documents per app that the proposals put in the bucket (and in outliers)."""
    final = {t: p.get("merge_into", t) for t, p in proposals.items()}
    bucket_ids = {t for t, p in proposals.items() if p.get("name") == bucket}
    in_bucket = np.array([t != -1 and final[int(t)] in bucket_ids for t in topics])
    out = {}
    for a in sorted(set(apps)) + ["all"]:
        m = np.ones(len(apps), bool) if a == "all" else apps == a
        out[a] = {"unspecified": round(float(in_bucket[m].mean()), 4), "outliers": round(float((topics[m] == -1).mean()), 4),
                  "unspecified_or_outlier": round(float((in_bucket[m] | (topics[m] == -1)).mean()), 4)}
    return out


# ---------------------------------------------------------------- modes

def fit_docs(tcfg: dict[str, Any]) -> pd.DataFrame:
    """The 60,014 fit documents in 6a order (unique, non-short PRIMARY complaints)."""
    comp = tp.load_complaints(tcfg)
    return comp[comp["in_fit"]].reset_index(drop=True)


def reduce_umap(emb: np.ndarray, tcfg: dict[str, Any]) -> np.ndarray:
    """Seeded UMAP exactly as BERTopic runs it (fit, then transform of the same data)."""
    from umap import UMAP

    u = tcfg["umap"]
    model = UMAP(n_neighbors=u["n_neighbors"], n_components=u["n_components"], min_dist=u["min_dist"],
                 metric=u["metric"], random_state=u["random_state"], low_memory=False)
    model.fit(emb)
    return model.transform(emb)


def compare(cfg: dict[str, Any], tcfg: dict[str, Any], fcfg: dict[str, Any]) -> None:
    """Per embedding model: UMAP once, HDBSCAN over the mcs grid, then BERTopic at the grid value whose
    topic count is closest to target_topics; compare NPMI, outliers and writing-style topics."""
    from bertopic import BERTopic
    from bertopic.dimensionality import BaseDimensionalityReduction
    from hdbscan import HDBSCAN

    fd = fit_docs(tcfg)
    docs = fd[tcfg["text_col"]].tolist()
    stop = stopwords_v2(tcfg, fcfg)
    h = tcfg["hdbscan"]
    results = {}
    for spec in cfg["embeddings"]:
        emb = embed(docs, spec, cfg)
        t0 = time.time()
        red = reduce_umap(emb, tcfg)
        logger.info("%s: UMAP %.0fs", spec["name"], time.time() - t0)
        grid = {}
        for mcs in cfg["mcs_grid"]:
            lab = HDBSCAN(min_cluster_size=mcs, metric=h["metric"], cluster_selection_method=h["cluster_selection_method"]).fit(red).labels_
            sizes = pd.Series(lab[lab != -1]).value_counts()
            grid[mcs] = {"n_topics": int(sizes.size), "outlier_share": float((lab == -1).mean()),
                         "largest_topic_share": float(sizes.max() / len(lab)) if sizes.size else 1.0}
        print(f"\n== {spec['name']}: grid " + "  ".join(f"mcs={m}: {g['n_topics']} topics, {g['outlier_share']:.0%} out, "
                                                         f"largest {g['largest_topic_share']:.0%}" for m, g in grid.items()))
        usable = {m: g for m, g in grid.items() if g["n_topics"] >= 5}
        r: dict[str, Any] = {"model": spec["model"], "prefix": spec["prefix"], "license": spec["license"],
                             "dim": int(emb.shape[1]), "grid": grid}
        if not usable:
            r["degenerate"] = "fewer than 5 topics at every grid value"
            results[spec["name"]] = r
            print("   degenerate: fewer than 5 topics at every min_cluster_size")
            continue
        mcs = min(usable, key=lambda m: (abs(usable[m]["n_topics"] - cfg["target_topics"]), -m))
        model = build(tcfg, fcfg, cfg, spec, mcs)
        model.umap_model = BaseDimensionalityReduction()  # reuse the reduction computed above
        model.hdbscan_model = HDBSCAN(min_cluster_size=mcs, metric=h["metric"],
                                      cluster_selection_method=h["cluster_selection_method"])
        model.fit(docs, embeddings=red)
        r.update({"min_cluster_size": mcs, **summarize(model, docs, cfg, stop, tcfg)})
        results[spec["name"]] = r
        print(f"   compared at mcs={mcs}: {r['n_topics']} topics, outliers {r['outlier_share']:.1%}, NPMI mean "
              f"{r['npmi_mean']:.3f} (size-weighted {r['npmi_size_weighted']:.3f}), style topics {r['n_style_topics']} "
              f"({r['style_docs_share']:.1%} of docs), largest {r['largest_topic_share']:.1%}")
        for t, d in r["topics"].items():
            flag = " [STYLE]" if t in r["style_topics"] else ""
            print(f"   T{t:<3} {d['size']:>6}  npmi {d['npmi']:+.2f}  {', '.join(d['keywords'][:8])}{flag}")
    out = Path(cfg["comparison_path"])
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps({"seed": cfg["seed"], "n_docs": len(docs), "mcs_grid": cfg["mcs_grid"],
                               "target_topics": cfg["target_topics"], "max_seq_length": cfg["max_seq_length"],
                               "style_threshold": cfg["style_threshold"],
                               "npmi": "mean pairwise NPMI of top-10 keywords, documents as windows; no co-occurrence = -1",
                               "results": results}, indent=2, ensure_ascii=False), encoding="utf-8")
    logger.info("wrote %s", out)


def gap_check(docs: list[str], topics: np.ndarray, names: dict[str, list[str]]) -> dict[str, Any]:
    """For each diagnostic gap: documents mentioning its terms and the topics they land in."""
    import re

    low = pd.Series(docs).str.lower()
    out = {}
    for gap, terms in names.items():
        pat = re.compile(r"\b(?:" + "|".join(re.escape(t) for t in terms) + r")\b")  # whole words ("lag" != "lagi")
        m = low.str.contains(pat).to_numpy()
        dist = pd.Series(topics[m]).value_counts()
        out[gap] = {"docs": int(m.sum()), "top_topics": {int(t): int(n) for t, n in dist.head(6).items()},
                    "share_in_top_topic": float(dist.iloc[0] / m.sum()) if m.any() else 0.0}
    return out


def fit_final(name: str, cfg: dict[str, Any], tcfg: dict[str, Any], fcfg: dict[str, Any]) -> None:
    """Refit with the chosen embedding, calibrate outlier reduction, save, and print reading samples."""
    spec = next(s for s in cfg["embeddings"] if s["name"] == name)
    comp = json.loads(Path(cfg["comparison_path"]).read_text(encoding="utf-8"))["results"][name]
    mcs = comp["min_cluster_size"]
    fd = fit_docs(tcfg)
    docs = fd[tcfg["text_col"]].tolist()
    neutral_info = None
    if cfg["neutralize_on"]:
        emb_input, neutral_info = neutralize_all(docs, cfg["neutralize"])
        logger.info("neutralized %d docs (%.1f%%): %s", neutral_info["docs_changed"],
                    100 * neutral_info["docs_changed_share"], neutral_info["replacements"])
    else:
        emb_input = docs
    emb = embed(emb_input, spec, cfg)
    model = build(tcfg, fcfg, cfg, spec, mcs)
    raw = np.asarray(model.fit_transform(docs, embeddings=emb)[0])

    # Calibrate the reduce_outliers threshold so ~target_share of documents stay outliers
    sim = tp.outlier_similarity(model, emb, raw)
    n_target = int(round(cfg["outliers"]["target_share"] * len(docs)))
    keep = min(n_target, len(sim))
    threshold = float(np.sort(sim)[keep - 1]) + 1e-9 if keep > 0 else 0.0  # the `keep` least similar stay outliers
    topics = np.asarray(model.reduce_outliers(docs, raw.tolist(), strategy=cfg["outliers"]["strategy"],
                                              embeddings=emb, threshold=threshold))
    model.update_topics(docs, topics=topics.tolist(), vectorizer_model=model.vectorizer_model, top_n_words=cfg["top_n_words"])
    logger.info("mcs=%d: %d topics (comparison had %d); outliers %.1f%% -> %.1f%% (threshold %.4f)", mcs,
                len(set(raw) - {-1}), comp["n_topics"], 100 * (raw == -1).mean(), 100 * (topics == -1).mean(), threshold)

    out_dir = Path(cfg["model_dir"])
    out_dir.mkdir(parents=True, exist_ok=True)
    model.save(str(out_dir / f"bertopic_{cfg['version']}.pkl"), serialization="pickle", save_embedding_model=False)
    pd.DataFrame({"reviewId": fd["reviewId"], "app": fd["app"], "topic_raw": raw, "topic": topics}).to_parquet(
        out_dir / f"fit_topics_{cfg['version']}.parquet", index=False)

    # Reading sample: 20 random documents per topic, excluding the 60 diagnostic rows
    # and every row read in an earlier version
    rd = cfg["reading"]
    excluded = set(pd.read_parquet(rd["diagnostic_key"])["reviewId"])
    for sp in cfg["exclude_samples"]:
        excluded |= set(pd.read_parquet(sp)["reviewId"])
    frame = fd.assign(topic=topics)
    frame = frame[~frame["reviewId"].isin(excluded)]
    sample = (frame[frame["topic"] != -1].groupby("topic").sample(n=rd["per_topic"], random_state=cfg["seed"])
              .sort_values("topic", kind="stable"))
    Path(cfg["sample_path"]).parent.mkdir(parents=True, exist_ok=True)
    sample[["reviewId", "app", "topic"]].to_parquet(cfg["sample_path"], index=False)

    stop = stopwords_v2(tcfg, fcfg)
    summ = summarize(model, docs, cfg, stop, tcfg)
    gaps = gap_check(docs, topics, cfg["gap_terms"])
    apps = tcfg["apps"]
    app_tot = fd["app"].value_counts()
    profile = app_profile(fd["app"].to_numpy(), topics)
    rep_path = Path(cfg["report_path"])
    rep = json.loads(rep_path.read_text(encoding="utf-8")) if rep_path.exists() else {"seed": cfg["seed"]}
    rep[cfg["report_key"]] = {"version": cfg["version"], "name": name, "model": spec["model"], "min_cluster_size": mcs,
                              "threshold": threshold, "neutralization": neutral_info,
                              "outlier_share_raw": float((raw == -1).mean()), "outlier_share_after": float((topics == -1).mean()),
                              "after_reduction": {k: v for k, v in summ.items() if k != "topics"},
                              "gap_check": gaps, "app_profile": profile, "reading_sample": cfg["sample_path"],
                              "reading_sample_excludes": [rd["diagnostic_key"], *cfg["exclude_samples"]]}
    if cfg.get("compare_with"):
        prev = cfg["versions"][cfg["compare_with"]]
        pft = pd.read_parquet(Path(prev["model_dir"]) / f"fit_topics_{cfg['compare_with']}.parquet")
        assert (pft["reviewId"].to_numpy() == fd["reviewId"].to_numpy()).all()
        rep[cfg["report_key"]]["before"] = {"version": cfg["compare_with"],
                                            "app_profile": app_profile(pft["app"].to_numpy(), pft["topic"].to_numpy())}
    rep_path.parent.mkdir(parents=True, exist_ok=True)
    rep_path.write_text(json.dumps(rep, indent=2, ensure_ascii=False, default=str), encoding="utf-8")

    print(f"\n== {cfg['version']} ({name}): {summ['n_topics']} topics; outliers {(raw == -1).mean():.1%} -> "
          f"{(topics == -1).mean():.1%}; NPMI {summ['npmi_mean']:.3f}; style topics {summ['style_topics']}")
    if neutral_info:
        print("neutralization:", json.dumps({k: v for k, v in neutral_info.items() if k != "examples"}))
    print("gap check:", json.dumps(gaps))
    print("app-dominated topics (>= 80% one app):", json.dumps(profile["app_dominated"]))
    if cfg.get("compare_with"):
        print(f"before ({cfg['compare_with']}):", json.dumps(rep[cfg["report_key"]]["before"]["app_profile"]["app_dominated"]))
    for t in sorted(set(topics) - {-1}):
        idx = topics == t
        shares = "  ".join(f"{a[:4]} {(fd['app'][idx] == a).sum() / app_tot[a]:.1%}" for a in apps)
        print(f"\n### T{t}  size {idx.sum()}  {shares}\n    kw: {', '.join(tp.keywords(model, t, 10))}")
        for i, d in enumerate(sample.loc[sample["topic"] == t, tcfg["text_col"]], 1):
            print(f"    {i:>2}. {d[:160]}")


def export(cfg: dict[str, Any], tcfg: dict[str, Any]) -> None:
    """topics_to_name_v2.xlsx: proposals + empty columns for confirmation, and the reading sample."""
    from bertopic import BERTopic
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    path = Path(cfg["xlsx_path"])
    if path.exists():
        raise FileExistsError(f"{path} exists; refusing to overwrite (it may contain your confirmations)")
    props = load_config(Path(cfg["proposals_path"]))["topics"]
    out_dir = Path(cfg["model_dir"])
    model = BERTopic.load(str(out_dir / f"bertopic_{cfg['version']}.pkl"))
    ft = pd.read_parquet(out_dir / f"fit_topics_{cfg['version']}.parquet")
    ids = sorted(int(t) for t in ft["topic"].unique() if t != -1)
    if sorted(int(k) for k in props) != ids:
        raise ValueError(f"proposals cover {sorted(props)}, model has {ids}")
    for t, p in props.items():
        if p.get("merge_into") is not None and p.get("merge_into") not in props:
            raise ValueError(f"T{t}: merge target {p['merge_into']} does not exist")
        if p.get("merge_into") is not None and props[p["merge_into"]].get("merge_into") is not None:
            raise ValueError(f"T{t}: merge target {p['merge_into']} is itself merged (chain)")
    apps = tcfg["apps"]
    tot = ft["app"].value_counts()
    rows = []
    for t in [-1] + ids:
        g = ft[ft["topic"] == t]
        p = props.get(t, {})
        rows.append({"topic_id": t, "size": len(g), **{f"share_{a}": (g["app"] == a).sum() / tot[a] for a in apps},
                     "keywords": ", ".join(tp.keywords(model, t, 10)) if t != -1 else "",
                     "proposed_name": p.get("name", "(outliers: not named)" if t == -1 else None),
                     "purity": f"{p['purity']}/20" if "purity" in p else None,
                     "proposed_merge_into": p.get("merge_into"), "reason": p.get("note"),
                     "topic_name": None, "merge_into": None, "notes": None})
    df = pd.DataFrame(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "topics"
    ws.append(list(df.columns))
    for r in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in r])
    widths = {"topic_id": 9, "size": 8, "keywords": 42, "proposed_name": 30, "purity": 8, "proposed_merge_into": 11,
              "reason": 50, "topic_name": 28, "merge_into": 11, "notes": 30}
    for i, col in enumerate(df.columns, 1):
        letter = ws.cell(1, i).column_letter
        ws.column_dimensions[letter].width = widths.get(col, 10)
        ws.cell(1, i).font = Font(bold=True)
        for row in range(2, len(df) + 2):
            c = ws.cell(row, i)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if col.startswith("share_"):
                c.number_format = "0.0%"
        if col in ("topic_name", "merge_into", "notes"):
            ws.cell(1, i).fill = PatternFill("solid", fgColor="FFF6D5")
    gray = PatternFill("solid", fgColor="E4E3DF")
    for ci in range(1, len(df.columns) + 1):
        ws.cell(2, ci).fill = gray
    mi = ws.cell(1, list(df.columns).index("merge_into") + 1).column_letter
    dv = DataValidation(type="list", formula1='"' + ",".join(map(str, ids)) + '"', allow_blank=True,
                        showErrorMessage=True, error="merge_into must be an existing topic_id")
    ws.add_data_validation(dv)
    dv.add(f"{mi}3:{mi}{len(df) + 1}")
    ws.freeze_panes = "C2"

    # Sheet 2: the documents read to propose names (purity is counted on these)
    texts = pd.read_parquet(tcfg["input_path"], columns=["reviewId", tcfg["text_col"]])
    smp = pd.read_parquet(cfg["sample_path"]).merge(texts, on="reviewId", how="left")
    ws2 = wb.create_sheet("reading_sample")
    ws2.append(["topic_id", "app", tcfg["text_col"]])
    for r in smp.itertuples(index=False):
        ws2.append([int(r.topic), r.app, getattr(r, tcfg["text_col"])])
    for col, w in zip("ABC", (10, 11, 110)):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions

    info = wb.create_sheet("README")
    for line in [
        f"{cfg['version']} topics (Phase 6c-1). proposed_* columns are Claude's proposals from reading 20 random documents per topic "
        "(sheet reading_sample).",
        "purity = how many of those 20 match proposed_name. Below 12/20 -> 'Unspecified complaint' bucket.",
        "Merges are proposed only between topics describing the same problem; mixed topics are never merged into "
        "specific or outage topics.",
        "Fill in: topic_name (or copy proposed_name), merge_into (topic_id; blank = keep), notes. "
        "Exactly one of topic_name / merge_into per topic. Topic -1 = outliers, not named.",
        "share_<app>: the topic's share of that app's fitted complaint documents (outliers in the denominator).",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 120
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    logger.info("wrote %s (%d topics + outliers, %d sample rows)", path, len(ids), len(smp))

    # Unspecified-bucket share per app under these proposals (and the previous version's)
    rep_path = Path(cfg["report_path"])
    rep = json.loads(rep_path.read_text(encoding="utf-8")) if rep_path.exists() else {}
    shares = {cfg["version"]: unspecified_shares(ft["app"].to_numpy(), ft["topic"].to_numpy(), props)}
    if cfg.get("compare_with"):
        prev = cfg["versions"][cfg["compare_with"]]
        pft = pd.read_parquet(Path(prev["model_dir"]) / f"fit_topics_{cfg['compare_with']}.parquet")
        pprops = load_config(Path(prev["proposals_path"]))["topics"]
        shares[cfg["compare_with"]] = unspecified_shares(pft["app"].to_numpy(), pft["topic"].to_numpy(), pprops)
    rep.setdefault(cfg["report_key"], {})["unspecified_shares"] = shares
    rep[cfg["report_key"]]["proposals"] = {"named": sum("name" in p for p in props.values()),
                                           "merged": sum("merge_into" in p for p in props.values())}
    rep_path.write_text(json.dumps(rep, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    for v, sh in shares.items():
        print(f"{v}: " + "  ".join(f"{a} unspecified {d['unspecified']:.1%} (+outliers {d['unspecified_or_outlier']:.1%})"
                                    for a, d in sh.items()))


def main() -> None:
    """CLI entry point."""
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--config", type=Path, default=Path("configs/topics_v2.yaml"))
    g = parser.add_mutually_exclusive_group(required=True)
    g.add_argument("--compare", action="store_true")
    g.add_argument("--fit", metavar="NAME")
    g.add_argument("--export", action="store_true")
    parser.add_argument("--version", default="v2", choices=["v2", "v3"])
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")
    for noisy in ("httpx", "huggingface_hub", "BERTopic", "numba", "sentence_transformers"):
        logging.getLogger(noisy).setLevel(logging.WARNING)

    cfg = load_config(args.config)
    tcfg = load_config(Path(cfg["topics_config"]))
    fcfg = load_config(Path(cfg["final_config"]))
    ver = cfg["versions"][args.version]
    cfg.update({k: v for k, v in ver.items() if k != "neutralize"}, version=args.version, neutralize_on=ver["neutralize"])
    np.random.seed(cfg["seed"])
    logger.info("seed=%d mode=%s", cfg["seed"], "compare" if args.compare else ("fit " + args.fit if args.fit else "export"))
    if args.compare:
        compare(cfg, tcfg, fcfg)
    elif args.fit:
        fit_final(args.fit, cfg, tcfg, fcfg)
    else:
        export(cfg, tcfg)


if __name__ == "__main__":
    main()
