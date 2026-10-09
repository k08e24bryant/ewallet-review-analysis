"""Phase 6c-2: apply the reviewed v3 topic names/merges and export a fresh fit check.

--apply
    1. validate data/topics/topics_to_name_v3.xlsx (same rules as 6b)
    2. assign every complaint row that was not a v3 fit document (short reviews,
       duplicates) with transform on the brand-neutral embedding input, then the
       same outlier reduction as the v3 fit (threshold, centroid topic embeddings)
    3. merge + name topics in the model, refresh keywords; update
       topic_assignments.parquet: v3 -> topic / topic_name, 6b kept as topic_v1 /
       topic_name_v1; check that assignments change only by the merges
--fit-check
    60 fresh rows (per app: 10 fitted documents + 5 transform-assigned rows),
    named topics only, one row per distinct text, excluding the 60 diagnostic rows
    and the 1,200 rows Claude read (by reviewId and by identical text)

No gold or dev labels are read.

Usage:
    uv run python -m src.topics_v3 --apply
    uv run python -m src.topics_v3 --fit-check
"""

from __future__ import annotations

import argparse
import json
import logging
import sys
import time
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import yaml

from src import topics as tp
from src import topics_final as tf
from src import topics_v2 as v2

logger = logging.getLogger("topics_v3")


def load_config(path: Path) -> dict[str, Any]:
    """Load a YAML config."""
    with path.open(encoding="utf-8") as f:
        return yaml.safe_load(f)


def centroid_reduce(emb: np.ndarray, raw: np.ndarray, centroids: np.ndarray, ids: list[int], threshold: float) -> np.ndarray:
    """BERTopic's 'embeddings' outlier strategy with fixed centroids: an outlier moves to the most
    similar topic if the cosine similarity is >= threshold."""
    from sklearn.metrics.pairwise import cosine_similarity

    out = raw.copy()
    m = raw == -1
    if m.any():
        sim = cosine_similarity(emb[m], centroids)
        best = sim.argmax(axis=1)
        out[np.flatnonzero(m)] = np.where(sim.max(axis=1) >= threshold, np.asarray(ids)[best], -1)
    return out


def apply(cfg: dict[str, Any]) -> None:  # noqa: PLR0915 - linear pipeline
    """Validate, transform the remaining complaint rows, merge + name, update assignments."""
    from bertopic import BERTopic

    tcfg = load_config(Path(cfg["topics_config"]))
    fcfg = load_config(Path(cfg["final_config"]))
    vcfg = load_config(Path(cfg["v2_config"]))
    text = tcfg["text_col"]
    spec = next(s for s in vcfg["embeddings"] if s["name"] == cfg["embedding"])

    model = BERTopic.load(cfg["model_in"], embedding_model=spec["model"])
    names = pd.read_excel(cfg["workbook"], sheet_name=cfg["sheet"])[["topic_id", "topic_name", "merge_into", "notes"]]
    errors, merge, kept = tf.validate(names, {int(t) for t in set(model.topics_)})
    groups: dict[int, list[int]] = {}
    for s, d in merge.items():
        groups.setdefault(d, []).append(s)
    print(f"\n== validation: {len(names) - 1} topics, {len(merge)} merged into {len(groups)} targets, {len(kept)} named ==")
    for d, ss in sorted(groups.items()):
        print(f"   {d:>3} {kept.get(d, '?')!r:<46} <- {sorted(ss)}")
    if errors:
        print("\nVALIDATION FAILED")
        for e in errors:
            print(" -", e)
        sys.exit(1)
    print("validation passed")

    # ---- fit documents (6a order) and the v3 fit topics
    comp = tp.load_complaints(tcfg)
    fit_df = comp[comp["in_fit"]].reset_index(drop=True)
    other = comp[~comp["in_fit"]].reset_index(drop=True)
    ft = pd.read_parquet(cfg["fit_topics"])
    if not (ft["reviewId"].to_numpy() == fit_df["reviewId"].to_numpy()).all():
        raise RuntimeError("fit_topics_v3 is not in fit-document order")
    raw_fit, red_fit = ft["topic_raw"].to_numpy(), ft["topic"].to_numpy()
    if not np.array_equal(red_fit, np.asarray(model.topics_)):
        raise RuntimeError("v3 model topics differ from fit_topics_v3")
    docs = fit_df[text].tolist()

    # ---- brand-neutral embeddings (cached) for fit docs and the other complaint texts
    ncfg = vcfg["neutralize"]
    emb_fit = v2.embed(v2.neutralize_all(docs, ncfg)[0], spec, vcfg)
    o_texts = other[text].drop_duplicates().tolist()
    o_input, o_neutral = v2.neutralize_all(o_texts, ncfg)
    emb_o = v2.embed(o_input, spec, vcfg)

    # Outlier reduction exactly as in the v3 fit: centroids of the raw fit topics, same threshold
    threshold = json.loads(Path(cfg["v3_report"]).read_text(encoding="utf-8"))["fit"]["threshold"]
    ids = sorted(int(t) for t in set(raw_fit) if t != -1)
    centroids = np.vstack([emb_fit[raw_fit == t].mean(axis=0) for t in ids])
    check_fit = centroid_reduce(emb_fit, raw_fit, centroids, ids, threshold)
    reduction_reproduced = bool(np.array_equal(check_fit, red_fit))
    if not reduction_reproduced:
        raise RuntimeError(f"centroid reduction does not reproduce the v3 fit ({np.mean(check_fit == red_fit):.4f} agree)")

    t0 = time.time()
    o_raw = np.asarray(model.transform(o_texts, embeddings=emb_o)[0])
    o_red = centroid_reduce(emb_o, o_raw, centroids, ids, threshold)
    logger.info("transform: %d unique texts in %.0fs; outliers %.1f%% -> %.1f%%", len(o_texts), time.time() - t0,
                100 * (o_raw == -1).mean(), 100 * (o_red == -1).mean())

    rng = np.random.default_rng(cfg["seed"])
    chk = rng.choice(len(docs), size=min(cfg["transform_check_n"], len(docs)), replace=False)
    chk_raw = np.asarray(model.transform([docs[i] for i in chk], embeddings=emb_fit[chk])[0])
    transform_check = {"n": int(len(chk)), "agreement_with_fit_topic_raw": float(np.mean(chk_raw == raw_fit[chk]))}

    # ---- merge + name + refresh keywords (assignments unchanged by the refresh)
    mapping = tf.apply_merges(model, docs, merge)
    label_of = {mapping[t]: n for t, n in kept.items()} | {-1: cfg["outlier_label"]}
    tf.refresh_keywords(model, docs, tcfg, fcfg)
    model.set_topic_labels(label_of)

    # ---- assignments: v3 for every complaint row; 6b kept as *_v1
    assign = pd.read_parquet(cfg["assignments_path"])
    if "topic_v1" in assign.columns:  # re-run: restore the 6b columns first
        assign = assign.assign(topic=assign["topic_v1"], topic_name=assign["topic_name_v1"]).drop(
            columns=[c for c in assign.columns if c.endswith(("_v1", "_v3", "_v3_raw"))])
    a = assign.rename(columns={"topic": "topic_v1", "topic_name": "topic_name_v1"})
    fit_map = dict(zip(fit_df["reviewId"], red_fit))
    fit_raw_map = dict(zip(fit_df["reviewId"], raw_fit))
    o_map, o_raw_map = dict(zip(o_texts, o_red)), dict(zip(o_texts, o_raw))
    txt = comp.set_index("reviewId")[text]
    in_fit_ids = set(fit_df["reviewId"])
    is_fit = a["reviewId"].isin(in_fit_ids)
    a["topic_v3_raw"] = np.where(is_fit, a["reviewId"].map(fit_raw_map), a["reviewId"].map(txt).map(o_raw_map))
    a["topic_v3_unmerged"] = np.where(is_fit, a["reviewId"].map(fit_map), a["reviewId"].map(txt).map(o_map))
    if a[["topic_v3_raw", "topic_v3_unmerged"]].isna().any().any():
        raise RuntimeError("some complaint rows have no v3 topic")
    a["topic_v3_raw"] = a["topic_v3_raw"].astype(int)
    a["topic_v3_unmerged"] = a["topic_v3_unmerged"].astype(int)
    a["topic"] = a["topic_v3_unmerged"].map(mapping).astype(int)
    a["topic_name"] = a["topic"].map(label_of)

    fit_final = a.set_index("reviewId").loc[fit_df["reviewId"], "topic"].to_numpy()
    expected = a["topic_v3_unmerged"].map(lambda t: merge.get(t, t))
    pairs = pd.DataFrame({"target": expected, "final": a["topic"]}).drop_duplicates()
    check = {
        "rows_before": int(len(assign)), "rows_after": int(len(a)),
        "rows_assigned_by_transform": int((~is_fit).sum()),
        "transform_rows_match_assigned_by_transform_flag": bool((~is_fit == a["assigned_by_transform"]).all()),
        "outlier_reduction_reproduces_v3_fit": reduction_reproduced,
        "model_topics_match_fit_rows": bool(np.array_equal(fit_final, np.asarray(model.topics_))),
        "partition_equals_requested_merges": bool(not pairs["target"].duplicated().any() and not pairs["final"].duplicated().any()),
        "outlier_rows_unchanged_by_merge": bool(((a["topic_v3_unmerged"] == -1) == (a["topic"] == -1)).all()),
        "rows_moved_by_merges": int(a["topic_v3_unmerged"].isin(list(merge)).sum()),
    }
    bad = [k for k, v in check.items() if isinstance(v, bool) and not v]
    if bad or check["rows_before"] != check["rows_after"]:
        raise RuntimeError(f"assignment check failed: {bad} {check}")

    cols = ["reviewId", "app", "topic", "topic_name", "topic_v3_unmerged", "topic_v3_raw", "topic_v1", "topic_name_v1",
            "topic_6a", "topic_before_outlier_reduction", "assigned_by_transform", "in_fit", "in_model_pool", "is_short"]
    a = a[cols]

    # ---- shares (unique texts and all rows) per app
    apps = tcfg["apps"]
    shares = {}
    for basis, part in (("unique_texts", a[a["in_model_pool"]]), ("all_rows", a)):
        ct = pd.crosstab(part["topic_name"], part["app"]).reindex(columns=apps, fill_value=0)
        ct["all"] = ct.sum(axis=1)
        shares[basis] = (ct / ct.sum(axis=0)).round(4).to_dict("index")
    kw = {label_of[t]: tp.keywords(model, t, 10) for t in sorted(set(model.topics_)) if t != -1}

    a.to_parquet(cfg["assignments_path"], index=False)
    model.save(cfg["model_out"], serialization="pickle", save_embedding_model=False)
    report = {
        "seed": cfg["seed"], "labels_used": "none (user-confirmed topic names; PRIMARY predictions; no gold or dev)",
        "workbook": cfg["workbook"],
        "validation": {"n_topics_v3": len(names) - 1, "n_named": len(kept), "merged": {int(d): sorted(s) for d, s in groups.items()}},
        "id_mapping_v3_to_final": {int(k): int(v) for k, v in mapping.items()},
        "outlier_reduction": {"threshold": threshold, "method": "centroids of raw v3 fit topics (brand-neutral embeddings)",
                              "transform_rows_outliers_before": float((o_raw == -1).mean()),
                              "transform_rows_outliers_after": float((o_red == -1).mean())},
        "transform_input_neutralization": {k: v for k, v in o_neutral.items() if k != "examples"},
        "transform_check_on_fit_docs": transform_check,
        "assignment_check": check,
        "outlier_share_all_rows": float((a["topic"] == -1).mean()),
        "shares": shares, "keywords": kw,
        "columns": {"topic/topic_name": "v3 (this phase)", "topic_v1/topic_name_v1": "6b topics (kept)",
                    "topic_v3_unmerged": "v3 before merges", "topic_v3_raw": "v3 before outlier reduction"},
    }
    rp = Path(cfg["report_path"])
    rp.parent.mkdir(parents=True, exist_ok=True)
    rp.write_text(json.dumps(report, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    logger.info("wrote %s, %s, %s", cfg["assignments_path"], cfg["model_out"], rp)

    print(f"\n== applied: {len(kept)} named topics + outliers; check: {check}")
    print("transform check:", transform_check, "| outliers all rows:", f"{report['outlier_share_all_rows']:.1%}")
    u = shares["unique_texts"]
    print(f"{'topic (unique texts)':<48}" + "".join(f"{c:>10}" for c in apps + ["all"]))
    for n in sorted(u, key=lambda n: -u[n]["all"]):
        print(f"{n[:47]:<48}" + "".join(f"{u[n][c]:>10.1%}" for c in apps + ["all"]))


def fit_check(cfg: dict[str, Any]) -> None:
    """Fresh 60-row fit-check workbook (stratified by app and assigned_by_transform) + hidden key."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    tcfg = load_config(Path(cfg["topics_config"]))
    vcfg = load_config(Path(cfg["v2_config"]))
    fc = cfg["fit_check"]
    if not all(isinstance(o, str) for o in fc["options"]):
        raise ValueError(f"fit_check.options must be strings: {fc['options']}")
    xlsx, key_path = Path(fc["xlsx_path"]), Path(fc["key_path"])
    for p in (xlsx, key_path):
        if p.exists():
            raise FileExistsError(f"{p} exists; refusing to overwrite (it may contain your judgments)")
    text = tcfg["text_col"]
    a = pd.read_parquet(cfg["assignments_path"])
    if "topic_v1" not in a.columns:
        raise RuntimeError("run --apply first")
    texts = pd.read_parquet(tcfg["input_path"], columns=["reviewId", text])
    a = a.merge(texts, on="reviewId", how="left", validate="one_to_one")

    excluded_ids: set[str] = set()
    for k in fc["exclude_keys"] + fc["exclude_reading"]:
        excluded_ids |= set(pd.read_parquet(k)["reviewId"])
    excluded_texts = set(a.loc[a["reviewId"].isin(excluded_ids), text])
    pool = a[(a["topic"] != -1) & ~a["reviewId"].isin(excluded_ids) & ~a[text].isin(excluded_texts)]
    pool = pool.sample(frac=1, random_state=cfg["seed"]).drop_duplicates(text)  # one random row per distinct text

    parts, avail = [], {}
    for app in tcfg["apps"]:
        for stratum, flag in (("fitted", False), ("transform", True)):
            cand = pool[(pool["app"] == app) & (pool["assigned_by_transform"] == flag)]
            n = fc["per_app"][stratum]
            avail[f"{app}|{stratum}"] = int(len(cand))
            if len(cand) < n:
                raise RuntimeError(f"{app}|{stratum}: only {len(cand)} candidates for {n}")
            parts.append(cand.sample(n=n, random_state=cfg["seed"]).assign(stratum=stratum))
    pick = pd.concat(parts).sample(frac=1, random_state=cfg["seed"]).reset_index(drop=True)
    pick.insert(0, "check_id", [f"{fc['id_prefix']}{i:0{fc['id_width']}d}" for i in range(1, len(pick) + 1)])
    assert not pick["reviewId"].isin(excluded_ids).any() and not pick[text].isin(excluded_texts).any()
    assert pick[text].is_unique

    wb = Workbook()
    ws = wb.active
    ws.title = "fit_check"
    ws.append(["check_id", "app", text, "topic_name", "fits", "notes"])
    for r in pick.itertuples():
        ws.append([r.check_id, r.app, getattr(r, text), r.topic_name, None, None])
    for col, width in zip("ABCDEF", (10, 11, 80, 38, 9, 30)):
        ws.column_dimensions[col].width = width
    for c in ws[1]:
        c.font = Font(bold=True)
    for col in ("E", "F"):
        ws[f"{col}1"].fill = PatternFill("solid", fgColor="FFF6D5")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"' + ",".join(fc["options"]) + '"', allow_blank=True,
                        showErrorMessage=True, error="Choose " + " / ".join(fc["options"]))
    ws.add_data_validation(dv)
    dv.add(f"E2:E{len(pick) + 1}")
    ws.freeze_panes = "A2"
    info = wb.create_sheet("README")
    for line in [
        "Does the topic_name describe what the review complains about? (fresh sample for the v3 topics)",
        "fits: yes = the topic is the main complaint; partly = related, or one of several complaints; no = wrong topic.",
        *[f"Rule: {r}" for r in vcfg["fit_check_rules"]],
        "Sample: 15 rows per app (10 fitted documents + 5 short/duplicate rows assigned by transform), seed 42, "
        "named topics only (outliers excluded), one row per distinct text.",
        "Excluded: the 60 rows of the first fit check and the 1,200 rows Claude read to name topics (by reviewId and text).",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 130
    xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx)
    key = pick[["check_id", "reviewId", "app", "stratum", "topic", "topic_name", "topic_v1", "topic_name_v1",
                "assigned_by_transform", "in_fit", "is_short", "in_model_pool"]]
    key.to_parquet(key_path, index=False)

    rp = Path(cfg["report_path"])
    rep = json.loads(rp.read_text(encoding="utf-8"))
    rep["fresh_fit_check"] = {
        "xlsx": str(xlsx), "key": str(key_path), "n": int(len(pick)),
        "design": {"per_app": fc["per_app"], "population": "complaint rows with a named v3 topic, one row per distinct text",
                   "excluded_rows": len(excluded_ids), "excluded_texts": len(excluded_texts), "available_per_stratum": avail},
        "per_topic": pick["topic_name"].value_counts().to_dict(),
        "decision_rule": cfg["decision"],
        "readme_rules": vcfg["fit_check_rules"],
    }
    rp.write_text(json.dumps(rep, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    logger.info("wrote %s (%d rows) and %s; excluded %d rows / %d texts", xlsx, len(pick), key_path,
                len(excluded_ids), len(excluded_texts))
    print(pd.crosstab(pick["app"], pick["stratum"]).to_string())
    print("topics covered:", pick["topic_name"].nunique(), pick["topic_name"].value_counts().to_dict())


def fit_rate(cfg: dict[str, Any]) -> None:  # noqa: PLR0915 - linear report
    """Score the filled fresh fit check, compare with 6b, and apply the pre-registered decision rule."""
    tcfg = load_config(Path(cfg["topics_config"]))
    fc, dec = cfg["fit_check"], cfg["decision"]
    options, apps = fc["options"], tcfg["apps"]
    ci = 0.95
    sheet = pd.read_excel(fc["xlsx_path"], sheet_name="fit_check")
    key = pd.read_parquet(fc["key_path"])
    errors = []
    if sheet["check_id"].tolist() != key["check_id"].tolist():
        errors.append("check_id order/content differs from the key")
    if (sheet["topic_name"].to_numpy() != key["topic_name"].to_numpy()).any():
        errors.append("topic_name column was changed")
    sheet["fits"] = sheet["fits"].map(lambda v: v.strip().lower() if isinstance(v, str) and v.strip() else None)
    if sheet["fits"].isna().any():
        errors.append(f"rows without a fits value: {sheet.loc[sheet['fits'].isna(), 'check_id'].tolist()}")
    bad = sheet[sheet["fits"].notna() & ~sheet["fits"].isin(options)]
    if len(bad):
        errors.append(f"unknown fits values: {bad[['check_id', 'fits']].values.tolist()}")
    if errors:
        raise ValueError("fit check failed validation: " + "; ".join(errors))
    pred = pd.read_parquet(tcfg["predictions_path"], columns=["reviewId", "score", "flagged_by"])
    df = (key.merge(sheet[["check_id", "fits", "notes", "text_clean"]], on="check_id", validate="one_to_one")
          .merge(pred, on="reviewId", how="left"))
    df["notes"] = df["notes"].map(lambda v: v.strip() if isinstance(v, str) and v.strip() else "")

    # Weighted all-rows rate: app x stratum cells weighted by their share of complaint rows with a
    # named topic (all rows, duplicates included, as used in Phase 7)
    a = pd.read_parquet(cfg["assignments_path"], columns=["app", "topic", "assigned_by_transform"])
    a = a[a["topic"] != -1].assign(stratum=lambda d: np.where(d["assigned_by_transform"], "transform", "fitted"))
    w = a.groupby(["app", "stratum"]).size() / len(a)
    cells = {c: g for c, g in df.groupby(["app", "stratum"])}

    def weighted(cs: dict, o: str) -> float:
        return float(sum(w[c] * (g["fits"] == o).mean() for c, g in cs.items()))

    rng = np.random.default_rng(cfg["seed"])
    boot = {o: [] for o in options}
    for _ in range(2000):
        res = {c: g.iloc[rng.integers(0, len(g), len(g))] for c, g in cells.items()}
        for o in options:
            boot[o].append(weighted(res, o))
    q = [(1 - ci) / 2 * 100, (1 + ci) / 2 * 100]
    weighted_all = {"weights_cell_share_of_rows": {f"{a_}|{s}": round(float(v), 4) for (a_, s), v in w.items()},
                    "method": "cell-weighted mean (app x fitted/transform); 2000 stratified bootstrap draws",
                    **{o: {"share": weighted(cells, o), "ci95": [float(x) for x in np.percentile(boot[o], q)]} for o in options}}

    fitted = df[df["stratum"] == "fitted"]
    metric = fitted["fits"].eq("yes").mean()
    adopt = bool(metric > dec["baseline_yes"])
    six_b = json.loads(Path(load_config(Path(cfg["final_config"]))["report_path"]).read_text(encoding="utf-8"))["fit_check_results"]
    broad = ["Frequent errors and outages", "Money lost, missing or taken without consent"]
    non_complaint = df[df["notes"].str.lower().str.contains("not a complaint|no complaint|praise")]
    cols = ["check_id", "app", "stratum", "topic_name", "notes", "text_clean"]
    r = {
        "n": int(len(df)), "options": options,
        "decision": {"rule": f"adopt v3 if the fitted-document 'yes' rate > {dec['baseline_yes']:.1%} (6b)",
                     "metric": "fitted_docs_yes_rate", "value": float(metric),
                     "ci95": tf.wilson(int(fitted["fits"].eq("yes").sum()), len(fitted), ci),
                     "baseline": dec["baseline_yes"], "baseline_ci95": six_b["raw"]["overall"]["yes"]["ci95"],
                     "adopt_v3": adopt},
        "fitted_docs": tf.fit_shares(fitted, options, ci),
        "transform_rows": tf.fit_shares(df[df["stratum"] == "transform"], options, ci),
        "all_rows_raw": tf.fit_shares(df, options, ci),
        "all_rows_weighted": weighted_all,
        "per_app": {x: tf.fit_shares(df[df["app"] == x], options, ci) for x in apps},
        "per_app_fitted": {x: tf.fit_shares(fitted[fitted["app"] == x], options, ci) for x in apps},
        "broad_vs_specific_fitted": {
            "broad": tf.fit_shares(fitted[fitted["topic_name"].isin(broad)], options, ci),
            "specific": tf.fit_shares(fitted[~fitted["topic_name"].isin(broad + ["Unspecified complaint"])], options, ci)},
        "broad_topics": broad,
        "per_topic": df.groupby("topic_name")["fits"].value_counts().unstack(fill_value=0).reindex(columns=options, fill_value=0)
        .assign(n=lambda t: t.sum(axis=1)).sort_values("n", ascending=False).to_dict("index"),
        "per_topic_fitted": fitted.groupby("topic_name")["fits"].value_counts().unstack(fill_value=0)
        .reindex(columns=options, fill_value=0).assign(n=lambda t: t.sum(axis=1)).sort_values("n", ascending=False).to_dict("index"),
        "non_complaints": {"count": int(len(non_complaint)), "by_stratum": non_complaint["stratum"].value_counts().to_dict(),
                           "flagged_by": non_complaint["flagged_by"].value_counts().to_dict(),
                           "rows": non_complaint[cols + ["score", "flagged_by"]].to_dict("records"),
                           "note": "notes describing praise or no complaint (sentiment errors of PRIMARY)"},
        "no_rows": df[df["fits"] == "no"].sort_values("check_id")[cols].to_dict("records"),
        "comparison_6b": {"6b_fitted_mostly_n60": {o: six_b["raw"]["overall"][o] for o in options},
                          "6b_broad_no": six_b["raw"]["by_topic_group"]["broad"]["no"]["share"],
                          "6b_specific_no": six_b["raw"]["by_topic_group"]["specific"]["no"]["share"],
                          "6b_non_complaints": len(six_b["not_complaint"]["possible_untagged"]),
                          "note": "6b sample: 15 per app, 59 of 60 fitted documents, outliers excluded"},
    }
    rp = Path(cfg["report_path"])
    rep = json.loads(rp.read_text(encoding="utf-8"))
    rep["fresh_fit_check_results"] = r
    rp.write_text(json.dumps(rep, indent=2, ensure_ascii=False, default=str), encoding="utf-8")

    def line(label: str, s: dict[str, Any]) -> str:
        return f"   {label:<26} n={s['n']:>2}  " + "  ".join(
            f"{o}: {s[o]['share']:5.1%} [{s[o]['ci95'][0]:.0%}, {s[o]['ci95'][1]:.0%}]" for o in options)

    d = r["decision"]
    print(f"\n== DECISION: fitted-document yes = {d['value']:.1%} [{d['ci95'][0]:.0%}, {d['ci95'][1]:.0%}] vs 6b "
          f"{d['baseline']:.1%} [{d['baseline_ci95'][0]:.0%}, {d['baseline_ci95'][1]:.0%}] -> adopt v3: {adopt}")
    print("\n== fit rates (Wilson 95% CI) ==")
    print(line("fitted docs (decision)", r["fitted_docs"]))
    print(line("transform rows", r["transform_rows"]))
    print(line("all 60 rows (raw)", r["all_rows_raw"]))
    wa = r["all_rows_weighted"]
    print("   all rows weighted          " + "  ".join(f"{o}: {wa[o]['share']:5.1%} [{wa[o]['ci95'][0]:.0%}, {wa[o]['ci95'][1]:.0%}]"
                                                    for o in options) + "   (bootstrap CI)")
    for x in apps:
        print(line(f"{x} (all 15)", r["per_app"][x]))
    for x in apps:
        print(line(f"{x} (fitted 10)", r["per_app_fitted"][x]))
    for g, s in r["broad_vs_specific_fitted"].items():
        print(line(f"fitted, {g} topics", s))
    print("\nper topic (all rows):", json.dumps(r["per_topic"]))
    nc = r["non_complaints"]
    print(f"\nnon-complaints: {nc['count']} {nc['by_stratum']} flagged_by {nc['flagged_by']}")
    print(f"\n== 'no' rows ({len(r['no_rows'])}) ==")
    for x in r["no_rows"]:
        print(f"   {x['check_id']} {x['app']:<9} {x['stratum']:<9} [{x['topic_name']}] {x['notes']}\n        {x['text_clean'][:140]!r}")


def main() -> None:
    """CLI entry point."""
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--config", type=Path, default=Path("configs/topics_v3.yaml"))
    g = parser.add_mutually_exclusive_group(required=True)
    g.add_argument("--apply", action="store_true")
    g.add_argument("--fit-check", action="store_true")
    g.add_argument("--fit-rate", action="store_true", help="score the filled fresh fit check and apply the decision rule")
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")
    for noisy in ("httpx", "huggingface_hub", "BERTopic", "numba", "sentence_transformers"):
        logging.getLogger(noisy).setLevel(logging.WARNING)
    cfg = load_config(args.config)
    np.random.seed(cfg["seed"])
    mode = "apply" if args.apply else ("fit-check" if args.fit_check else "fit-rate")
    logger.info("seed=%d mode=%s", cfg["seed"], mode)
    if args.apply:
        apply(cfg)
    elif args.fit_check:
        fit_check(cfg)
    else:
        fit_rate(cfg)


if __name__ == "__main__":
    main()
