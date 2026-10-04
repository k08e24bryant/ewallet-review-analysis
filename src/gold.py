"""Phase 3b: import the manually labeled gold workbook.

Steps:
    1. validate the workbook against gold_key.parquet (labels, note tags, ids, texts)
    2. replace invalid gold rows with reserve rows from the same hidden stratum
       (app x weak_label), in R-id order; stop and list R-ids that still need a label

Usage:
    uv run python -m src.gold --config configs/gold.yaml
"""

from __future__ import annotations

import argparse
import logging
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import yaml

logger = logging.getLogger("gold")

EXIT_INVALID_WORKBOOK = 1
EXIT_NEEDS_RESERVE_LABELS = 2
EXIT_RESERVE_EXHAUSTED = 3


def load_config(path: Path) -> dict[str, Any]:
    """Load the gold YAML config."""
    with path.open(encoding="utf-8") as f:
        return yaml.safe_load(f)


def read_sheet(path: Path, sheet: str) -> pd.DataFrame:
    """Read one labeling sheet into a frame with gold_id, text_clean, label, notes."""
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0][:4]]
    expected = ["gold_id", "text_clean", "label", "notes"]
    if header != expected:
        raise ValueError(f"sheet {sheet!r} header {header} != {expected}")
    df = pd.DataFrame([r[:4] for r in rows[1:] if any(v is not None for v in r[:4])], columns=expected)
    df["label"] = df["label"].map(lambda v: v.strip().lower() if isinstance(v, str) and v.strip() else None)
    return df


def parse_tags(value: Any) -> list[str]:
    """Split a notes cell into lowercase tags (comma-separated); empty -> []."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return []
    return [t.strip().lower() for t in str(value).split(",") if t.strip()]


@dataclass
class Validation:
    """Result of workbook validation; ``errors`` empty means it passed."""

    errors: list[str] = field(default_factory=list)
    details: dict[str, Any] = field(default_factory=dict)


def validate(sheets: dict[str, pd.DataFrame], key: pd.DataFrame, pool: pd.DataFrame, cfg: dict[str, Any]) -> Validation:
    """Check labels, note tags, and that ids/texts/order match the key."""
    v = Validation()
    allowed = set(cfg["labels"]) | {cfg["invalid_label"]}
    tags_ok = set(cfg["note_tags"])
    text_by_id = pool.set_index("reviewId")["text_clean"]

    for name, df in sheets.items():
        k = key[key["set"] == name].reset_index(drop=True)
        if df["gold_id"].tolist() != k["gold_id"].tolist():
            missing = sorted(set(k["gold_id"]) - set(df["gold_id"]))
            extra = sorted(set(df["gold_id"]) - set(k["gold_id"]))
            v.errors.append(
                f"[{name}] gold_id order/content differs from key (missing={missing[:10]}, extra={extra[:10]})"
            )
            continue
        expected_text = k["reviewId"].map(text_by_id).tolist()
        changed = [gid for gid, a, b in zip(df["gold_id"], df["text_clean"], expected_text) if a != b]
        if changed:
            v.errors.append(f"[{name}] {len(changed)} texts differ from the exported text: {changed[:10]}")

        bad_label = df[df["label"].notna() & ~df["label"].isin(allowed)]
        if len(bad_label):
            v.errors.append(f"[{name}] unknown labels: {bad_label[['gold_id', 'label']].values.tolist()}")
        bad_tags = [
            (gid, t) for gid, notes in zip(df["gold_id"], df["notes"]) for t in parse_tags(notes) if t not in tags_ok
        ]
        if bad_tags:
            v.errors.append(f"[{name}] unknown note tags: {bad_tags[:20]}")

    gold = sheets["gold"]
    blank = gold[gold["label"].isna()]
    if len(blank):
        v.errors.append(f"[gold] {len(blank)} rows without a label: {blank['gold_id'].tolist()}")
        v.details["blank_gold"] = blank[["gold_id", "text_clean"]].to_dict("records")
    return v


@dataclass
class ReplacementPlan:
    """Which reserve row fills each invalid gold slot, and what is still missing."""

    filled: list[dict[str, Any]] = field(default_factory=list)          # gold_id -> reserve gold_id
    needs_label: list[str] = field(default_factory=list)                # reserve R-ids to label next
    exhausted: dict[str, int] = field(default_factory=dict)             # stratum -> slots that cannot be filled
    reserve_invalid: list[str] = field(default_factory=list)


def plan_replacements(gold: pd.DataFrame, reserve: pd.DataFrame, key: pd.DataFrame, invalid: str) -> ReplacementPlan:
    """Walk each stratum's reserve rows in id order (R-ids, then X-ids) to fill its invalid gold slots.

    An unlabeled reserve row that is needed stops the walk for that stratum (it
    must be labeled first). Invalid reserve rows are skipped.
    """
    plan = ReplacementPlan()
    k = key.set_index("gold_id")
    g = gold.join(k[["app", "weak_label"]], on="gold_id")
    r = reserve.join(k[["app", "weak_label"]], on="gold_id").sort_values("gold_id")

    for (app, wl), slots in g[g["label"] == invalid].groupby(["app", "weak_label"]):
        queue = r[(r["app"] == app) & (r["weak_label"] == wl)].itertuples()
        open_slots = list(slots["gold_id"])
        while open_slots:
            row = next(queue, None)
            if row is None:
                plan.exhausted[f"{app}|{wl}"] = len(open_slots)
                break
            if pd.isna(row.label):  # blank cells arrive as None or NaN
                # Needed but unlabeled: ask for this one and any further ones
                # this stratum would need if all of them turn out valid.
                plan.needs_label.append(row.gold_id)
                for _ in range(len(open_slots) - 1):
                    nxt = next(queue, None)
                    if nxt is None:
                        plan.exhausted[f"{app}|{wl}"] = plan.exhausted.get(f"{app}|{wl}", 0) + 1
                    elif pd.isna(nxt.label):
                        plan.needs_label.append(nxt.gold_id)
                break
            if row.label == invalid:
                plan.reserve_invalid.append(row.gold_id)
                continue
            plan.filled.append({"gold_id": open_slots.pop(0), "reserve_id": row.gold_id, "stratum": f"{app}|{wl}"})
    plan.needs_label.sort()
    return plan


def used_review_ids(key: pd.DataFrame, splits_dir: Path) -> dict[str, set[str]]:
    """reviewId sets for train, val, and every set in the gold key."""
    sets = {name: set(pd.read_parquet(splits_dir / f"{name}.parquet")["reviewId"]) for name in ("train", "val")}
    for name, g in key.groupby("set"):
        sets[name] = set(g["reviewId"])
    return sets


def draw_reserve_extra(
    pool: pd.DataFrame, key: pd.DataFrame, strata: list[str], cfg: dict[str, Any]
) -> pd.DataFrame:
    """Draw extra reserve rows for the given 'app|weak_label' strata from unused pool rows.

    Excludes train, val, and all gold-key reviewIds. Returns new key rows
    (set='reserve_extra') with X-ids assigned in a shuffled order across strata.
    """
    ex = cfg["reserve_extra"]
    used = set().union(*used_review_ids(key, Path(cfg["splits_dir"])).values())
    spare = pool[~pool["reviewId"].isin(used)]
    parts = []
    for stratum in sorted(strata):
        app, wl = stratum.split("|")
        rows = spare[(spare["app"] == app) & (spare["weak_label"] == wl)]
        if len(rows) < ex["per_stratum"]:
            raise RuntimeError(
                f"stratum {stratum} has only {len(rows)} unused pool rows (needs {ex['per_stratum']}); "
                "rows would have to come out of train/val, which needs a split rebuild"
            )
        parts.append(rows.sample(n=ex["per_stratum"], random_state=cfg["seed"]))
    drawn = pd.concat(parts, ignore_index=True)
    rng = np.random.default_rng(cfg["seed"])
    drawn = drawn.iloc[rng.permutation(len(drawn))].reset_index(drop=True)
    drawn.insert(0, "gold_id", [f"{ex['id_prefix']}{i:0{ex['id_width']}d}" for i in range(1, len(drawn) + 1)])
    drawn["set"] = "reserve_extra"
    return drawn


def export_reserve_extra(drawn: pd.DataFrame, key: pd.DataFrame, cfg: dict[str, Any], label_cfg: dict[str, Any]) -> pd.DataFrame:
    """Write reserve_extra.xlsx (refuses to overwrite) and append the rows to gold_key.parquet."""
    from openpyxl import Workbook

    from src.label import write_label_sheet

    ex = cfg["reserve_extra"]
    path = Path(ex["xlsx_path"])
    if path.exists():
        raise FileExistsError(f"{path} exists; refusing to overwrite (it may contain labels)")
    if (key["set"] == "reserve_extra").any():
        raise RuntimeError("gold_key.parquet already has reserve_extra rows")

    wb = Workbook()
    ws = wb.active
    ws.title = ex["sheet"]
    write_label_sheet(ws, drawn, label_cfg["gold"])
    wb.save(path)

    new_key = pd.concat([key, drawn[key.columns]], ignore_index=True)
    sets = used_review_ids(new_key, Path(cfg["splits_dir"]))
    names = list(sets)
    for i, a in enumerate(names):
        for b in names[i + 1 :]:
            overlap = len(sets[a] & sets[b])
            assert overlap == 0, f"{overlap} reviewIds overlap between {a} and {b}"
    new_key.to_parquet(cfg["key_path"], index=False)
    logger.info("wrote %s (%d rows) and appended them to %s; overlap checks passed for %s",
                path, len(drawn), cfg["key_path"], names)
    return new_key


def main() -> None:
    """CLI entry point."""
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--config", type=Path, default=Path("configs/gold.yaml"))
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")

    cfg = load_config(args.config)
    logger.info("seed=%s config=%s", cfg["seed"], args.config)
    label_cfg = load_config(Path(cfg["label_config"]))
    key = pd.read_parquet(cfg["key_path"])
    pool = pd.read_parquet(cfg["pool_path"])

    def load_sheets() -> dict[str, pd.DataFrame]:
        sheets = {name: read_sheet(Path(cfg["workbook"]), name) for name in ("gold", "reserve")}
        extra = cfg["reserve_extra"]
        if Path(extra["xlsx_path"]).exists():
            sheets["reserve_extra"] = read_sheet(Path(extra["xlsx_path"]), extra["sheet"])
        return sheets

    def reserve_rows(sheets: dict[str, pd.DataFrame]) -> pd.DataFrame:
        return pd.concat([sheets[s] for s in ("reserve", "reserve_extra") if s in sheets], ignore_index=True)

    sheets = load_sheets()
    v = validate(sheets, key, pool, cfg)
    gold = sheets["gold"]
    print_gold_summary(gold, key, cfg)

    plan = plan_replacements(gold, reserve_rows(sheets), key, cfg["invalid_label"])
    if v.errors:
        print("\nVALIDATION FAILED")
        for e in v.errors:
            print(" -", e)
        for row in v.details.get("blank_gold", []):
            print(f"   {row['gold_id']}: {row['text_clean']!r}")
        print_plan(plan, preview=True)
        sys.exit(EXIT_INVALID_WORKBOOK)
    print("\nvalidation passed:", {name: len(df) for name, df in sheets.items()})

    if plan.exhausted and "reserve_extra" not in sheets:
        logger.info("reserve runs out for %s; drawing reserve_extra", plan.exhausted)
        drawn = draw_reserve_extra(pool, key, list(plan.exhausted), cfg)
        key = export_reserve_extra(drawn, key, cfg, label_cfg)
        sheets = load_sheets()
        v = validate(sheets, key, pool, cfg)
        if v.errors:
            raise RuntimeError(f"reserve_extra failed validation right after export: {v.errors}")
        plan = plan_replacements(gold, reserve_rows(sheets), key, cfg["invalid_label"])
        extra = key[key["set"] == "reserve_extra"]
        print("\n== reserve_extra drawn ==")
        print(extra.groupby(["app", "weak_label"]).size().to_string())

    print_plan(plan, preview=False)
    if plan.exhausted:
        sys.exit(EXIT_RESERVE_EXHAUSTED)
    if plan.needs_label:
        sys.exit(EXIT_NEEDS_RESERVE_LABELS)

    finalize(sheets, key, pool, plan, cfg, label_cfg)


def finalize(
    sheets: dict[str, pd.DataFrame],
    key: pd.DataFrame,
    pool: pd.DataFrame,
    plan: ReplacementPlan,
    cfg: dict[str, Any],
    label_cfg: dict[str, Any],
) -> None:
    """Steps 3-5: save gold_labels.parquet, write the agreement report + figures, export the blind relabel set."""
    import json

    from src import gold_report as gr

    final = gr.assemble_final(sheets, key, plan.filled, cfg["invalid_label"])
    splits = Path(cfg["splits_dir"])
    checks = gr.check_final(
        final,
        cfg,
        set(pd.read_parquet(splits / "train.parquet")["reviewId"]),
        set(pd.read_parquet(splits / "val.parquet")["reviewId"]),
    )
    out_cols = ["gold_id", "slot_id", "reviewId", "app", "score", "weak_label", "gold_label", "notes", "source", "reserve_set"]
    final[out_cols].to_parquet(cfg["output_path"], index=False)
    logger.info("wrote %s: %s", cfg["output_path"], checks)

    weighted = gr.add_weights(final, pool)
    logger.info("computing agreement report (bootstrap n=%d)", cfg["bootstrap"]["n"])
    rep = gr.agreement_report(weighted, cfg)
    rep = {
        "seed": cfg["seed"],
        "checks": checks,
        "weighting": (
            "raw = the 600 gold rows as sampled (50 per app x weak_label); reweighted = each row weighted "
            "by in_model_pool stratum size / 50, estimating the pool of unique review texts. Quote "
            "reweighted numbers for population statements; raw describes the gold set itself."
        ),
        **rep,
    }
    fig_dir = Path(cfg["figures_dir"])
    fig_dir.mkdir(parents=True, exist_ok=True)
    gr.plot_confusion(rep, fig_dir / cfg["figures"]["confusion"])
    gr.plot_mismatch(rep, fig_dir / cfg["figures"]["mismatch"])
    Path(cfg["report_path"]).write_text(
        json.dumps(rep, indent=2, ensure_ascii=False, default=lambda o: o.item() if hasattr(o, "item") else str(o)),
        encoding="utf-8",
    )
    logger.info("wrote %s and figures in %s", cfg["report_path"], fig_dir)

    relabel_key = gr.export_relabel(final, pool, cfg, label_cfg)
    logger.info("wrote %s (%d rows) and %s", cfg["relabel"]["xlsx_path"], len(relabel_key), cfg["relabel"]["key_path"])
    print_final_report(rep, relabel_key)


def print_final_report(rep: dict[str, Any], relabel_key: pd.DataFrame) -> None:
    """Print the Phase 3b report tables."""
    print("\n== gold_label distribution (count | raw share | reweighted share) ==")
    for app, d in rep["gold_label_distribution"].items():
        cells = "  ".join(
            f"{l}: {d['counts'][l]:>3} | {d['raw_share'][l]:5.1%} | {d['reweighted_share'][l]:5.1%}" for l in d["counts"]
        )
        print(f"{app:<10} {cells}")

    print("\n== Cohen's kappa, weak vs gold (raw | reweighted [95% CI]) ==")
    for app, k in rep["kappa"].items():
        lo, hi = k["reweighted_ci95"]
        print(f"{app:<10} {k['raw']:.3f} | {k['reweighted']:.3f} [{lo:.3f}, {hi:.3f}]")

    cm = rep["confusion_matrix"]
    print("\n== confusion matrix (rows = weak, cols = gold) raw counts ==")
    print(pd.DataFrame(cm["raw_counts"], index=cm["labels"], columns=cm["labels"]).to_string())
    print("reweighted row share:")
    print(pd.DataFrame(cm["reweighted_row_share"], index=cm["labels"], columns=cm["labels"]).round(3).to_string())

    print("\n== weak label as a predictor of gold (precision / recall) ==")
    for mode in ("raw", "reweighted"):
        q = rep["weak_label_quality_vs_gold"][mode]
        print(f"{mode:<10} " + "  ".join(f"{l}: P={v['precision']:.2f} R={v['recall']:.2f}" for l, v in q.items()))

    print("\n== rating-text mismatch (raw | reweighted [95% CI]) ==")
    for app, m in rep["rating_text_mismatch"].items():
        a, b = m["pos_stars_labeled_negative"], m["neg_stars_labeled_positive"]
        print(
            f"{app:<10} 4-5* -> negative: {a['raw']:5.1%} | {a['reweighted']:5.1%} "
            f"[{a['reweighted_ci95'][0]:.1%}, {a['reweighted_ci95'][1]:.1%}]   "
            f"1-2* -> positive: {b['raw']:5.1%} | {b['reweighted']:5.1%} "
            f"[{b['reweighted_ci95'][0]:.1%}, {b['reweighted_ci95'][1]:.1%}]"
        )
    print("by star (reweighted):", {s: {k: round(v, 3) for k, v in d.items()} for s, d in rep["gold_label_by_star_reweighted"].items()})

    print("\nnotes tags:", rep["notes_tags"])
    print("replaced invalid per stratum:", {k: v["replaced"] for k, v in rep["replaced_invalid_per_stratum"].items()})
    print("relabel set per app:", relabel_key["app"].value_counts().to_dict())


def print_gold_summary(gold: pd.DataFrame, key: pd.DataFrame, cfg: dict[str, Any]) -> None:
    """Print labels per hidden stratum, note tags, and every invalid row's text."""
    k = key.set_index("gold_id")
    g = gold.join(k[["app", "weak_label"]], on="gold_id")
    print("\n== gold labels per hidden stratum ==")
    print(pd.crosstab([g["app"], g["weak_label"]], g["label"].fillna("<blank>"), margins=True).to_string())
    tags = pd.Series([t for n in gold["notes"] for t in parse_tags(n)], dtype="object").value_counts()
    print("note tags:", tags.to_dict())
    inv = g[g["label"] == cfg["invalid_label"]].sort_values(["app", "weak_label", "gold_id"])
    print(f"\n== invalid gold rows ({len(inv)}) ==")
    for r in inv.itertuples():
        tag = f"  [{r.notes}]" if isinstance(r.notes, str) and r.notes else ""
        print(f"  {r.app:<9} {r.weak_label:<8} {r.gold_id}  {r.text_clean!r}{tag}")


def print_plan(plan: ReplacementPlan, preview: bool) -> None:
    """Print the replacement plan (as a preview when validation failed)."""
    title = "replacement plan (PREVIEW: assumes blank gold rows are not invalid)" if preview else "replacement plan"
    print(f"\n== {title} ==")
    print("filled from labeled reserve:", len(plan.filled))
    if plan.reserve_invalid:
        print("reserve rows labeled invalid (skipped):", plan.reserve_invalid)
    if plan.needs_label:
        print(f"reserve rows you must label ({len(plan.needs_label)}): {', '.join(plan.needs_label)}")
    if plan.exhausted:
        print("strata that run out of reserve rows (slots still empty):", plan.exhausted)


if __name__ == "__main__":
    main()
