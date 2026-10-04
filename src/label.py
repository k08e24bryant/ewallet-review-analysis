"""Phase 3a: weak labels, gold sample export for manual labeling, and training splits.

Steps:
    1. weak_label from star score for rows in in_model_pool
    2. gold sample (app x weak_label strata) + reserve, exported to xlsx without
       any score/app/label hints; the hidden key goes to gold_key.parquet
    3. GUIDELINES.md for manual labeling, with examples from the pool
    4. train/val splits from the pool minus gold and reserve

Usage:
    uv run python -m src.label --config configs/label.yaml
"""

from __future__ import annotations

import argparse
import logging
import sys
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import yaml
from sklearn.model_selection import train_test_split

logger = logging.getLogger("label")


def load_config(path: Path) -> dict[str, Any]:
    """Load the label YAML config."""
    with path.open(encoding="utf-8") as f:
        return yaml.safe_load(f)


def build_pool(df: pd.DataFrame, weak_label_map: dict[int, str]) -> pd.DataFrame:
    """Return in_model_pool rows with a weak_label column derived from score."""
    pool = df[df["in_model_pool"]].copy()
    pool["weak_label"] = pool["score"].map({int(k): v for k, v in weak_label_map.items()})
    if pool["weak_label"].isna().any():
        raise ValueError(f"unmapped scores: {sorted(pool.loc[pool['weak_label'].isna(), 'score'].unique())}")
    return pool.reset_index(drop=True)


def sample_gold(pool: pd.DataFrame, cfg: dict[str, Any]) -> pd.DataFrame:
    """Sample per (app, weak_label) stratum: the first ``per_stratum`` rows are gold, the rest reserve.

    Returns a key frame with gold_id, set, reviewId, app, score, weak_label, text_clean,
    in a shuffled order so apps and classes are mixed.
    """
    g = cfg["gold"]
    n_total = g["per_stratum"] + g["reserve_per_stratum"]
    parts = []
    for (app, label), stratum in pool.groupby(["app", "weak_label"], sort=True):
        if len(stratum) < n_total:
            raise ValueError(f"stratum {app}/{label} has {len(stratum)} rows, needs {n_total}")
        s = stratum.sample(n=n_total, random_state=cfg["seed"])
        s = s.assign(set=["gold"] * g["per_stratum"] + ["reserve"] * g["reserve_per_stratum"])
        parts.append(s)
    picked = pd.concat(parts, ignore_index=True)

    rng = np.random.default_rng(cfg["seed"])
    key_parts = []
    for set_name, prefix in (("gold", "G"), ("reserve", "R")):
        s = picked[picked["set"] == set_name]
        s = s.iloc[rng.permutation(len(s))].reset_index(drop=True)
        s.insert(0, "gold_id", [f"{prefix}{i:04d}" for i in range(1, len(s) + 1)])
        key_parts.append(s)
    key = pd.concat(key_parts, ignore_index=True)
    return key[["gold_id", "set", "reviewId", "app", "score", "weak_label", "is_short", "text_clean"]]


def write_label_sheet(ws, rows: pd.DataFrame, cfg_gold: dict[str, Any]) -> None:
    """Fill one worksheet: header, text cells forced to string, dropdown, wrap, widths, frozen header."""
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    headers = ["gold_id", "text_clean", "label", "notes"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    for r, (gid, text) in enumerate(zip(rows["gold_id"], rows["text_clean"]), start=2):
        for c, value in enumerate([gid, text, None, None], start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, str):
                cell.data_type = "s"  # never let "=..."/"@..." become a formula
            cell.alignment = wrap

    options = ",".join(cfg_gold["label_options"])
    dv = DataValidation(
        type="list",
        formula1=f'"{options}"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Invalid label",
        error=f"Choose one of: {options}",
    )
    ws.add_data_validation(dv)
    dv.add(f"C2:C{len(rows) + 1}")

    for col_letter, name in zip("ABCD", headers):
        ws.column_dimensions[col_letter].width = cfg_gold["column_widths"][name]
    ws.freeze_panes = "A2"


def export_gold(key: pd.DataFrame, cfg: dict[str, Any]) -> None:
    """Write the labeling workbook (sheets: gold, reserve) and the hidden key parquet."""
    from openpyxl import Workbook

    g = cfg["gold"]
    wb = Workbook()
    ws_gold = wb.active
    ws_gold.title = "gold"
    write_label_sheet(ws_gold, key[key["set"] == "gold"], g)
    write_label_sheet(wb.create_sheet("reserve"), key[key["set"] == "reserve"], g)

    xlsx_path = Path(g["xlsx_path"])
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    if xlsx_path.exists():
        raise FileExistsError(f"{xlsx_path} exists; refusing to overwrite (it may contain manual labels)")
    wb.save(xlsx_path)
    key.drop(columns=["text_clean", "is_short"]).to_parquet(g["key_path"], index=False)
    logger.info("wrote %s and %s", xlsx_path, g["key_path"])


def allocate(counts: pd.Series, total: int, mode: str) -> pd.Series:
    """Split ``total`` samples across strata (largest remainder), never exceeding a stratum's size."""
    if total >= counts.sum():
        return counts.copy()
    if mode == "proportional":
        raw = counts / counts.sum() * total
    elif mode == "equal":
        raw = pd.Series(total / len(counts), index=counts.index)
    else:
        raise ValueError(f"unknown week_allocation {mode!r}")
    alloc = np.minimum(np.floor(raw), counts).astype(int)
    # Hand out what is left: largest remainder first, only to strata with room.
    while alloc.sum() < total:
        room = counts - alloc
        rem = (raw - alloc).where(room > 0, -np.inf)
        k = min(int(total - alloc.sum()), int((room > 0).sum()))
        alloc[rem.nlargest(k).index] += 1
    return alloc


def downsample_app(rows: pd.DataFrame, cfg_splits: dict[str, Any], seed: int) -> pd.DataFrame:
    """Down-sample one app's rows to ``per_app``, stratified by week."""
    by = cfg_splits["stratify_by"]
    counts = rows[by].value_counts().sort_index()
    alloc = allocate(counts, cfg_splits["per_app"], cfg_splits["week_allocation"])
    parts = [
        rows[rows[by] == week].sample(n=int(n), random_state=seed)
        for week, n in alloc.items()
        if n > 0
    ]
    return pd.concat(parts)


def build_splits(pool: pd.DataFrame, key: pd.DataFrame, cfg: dict[str, Any]) -> dict[str, pd.DataFrame]:
    """Exclude gold + reserve, down-sample per app, then stratified 90/10 train/val."""
    sp = cfg["splits"]
    candidates = pool[~pool["reviewId"].isin(key["reviewId"])]
    sampled = []
    for app, rows in candidates.groupby("app"):
        if app in sp["use_all_rows"]:
            sampled.append(rows)
        else:
            sampled.append(downsample_app(rows, sp, cfg["seed"]))
    data = pd.concat(sampled, ignore_index=True)
    strata = data["app"] + "|" + data["weak_label"]
    train, val = train_test_split(data, test_size=sp["val_frac"], stratify=strata, random_state=cfg["seed"])
    return {"train": train.reset_index(drop=True), "val": val.reset_index(drop=True)}


def assert_no_overlap(sets: dict[str, set[str]]) -> dict[str, int]:
    """Assert pairwise-disjoint reviewId sets; returns the pairwise overlap counts (all zero)."""
    names = list(sets)
    overlaps = {}
    for i, a in enumerate(names):
        for b in names[i + 1 :]:
            n = len(sets[a] & sets[b])
            overlaps[f"{a}&{b}"] = n
            assert n == 0, f"{n} reviewIds overlap between {a} and {b}"
    return overlaps


def render_guidelines(pool: pd.DataFrame, key: pd.DataFrame, cfg: dict[str, Any]) -> str:
    """Render GUIDELINES.md; examples come from config and must not be gold/reserve rows."""
    ex_cfg = cfg["guidelines"]["examples"]
    by_id = pool.set_index("reviewId")
    blocked = set(key["reviewId"])
    examples: dict[str, list[str]] = {label: [] for label in cfg["gold"]["label_options"]}
    for e in ex_cfg:
        rid = e["reviewId"]
        if rid in blocked:
            raise ValueError(f"guideline example {rid} is in the gold/reserve set")
        if rid not in by_id.index:
            raise ValueError(f"guideline example {rid} is not in the model pool")
        text = by_id.at[rid, "text_clean"].replace("|", "\\|")
        examples[e["label"]].append(f"| {text} | {e['note']} |")

    def table(label: str) -> str:
        rows = examples[label]
        if not rows:
            return "_(no examples configured)_"
        return "| Review | Why |\n|---|---|\n" + "\n".join(rows)

    g = cfg["gold"]
    n_gold = g["per_stratum"] * len(cfg["apps"]) * len(cfg["classes"])
    return f"""# Gold set labeling guidelines

You are labeling **{n_gold} reviews** (sheet `gold`) of Indonesian e-wallet
apps by **text only**. Star ratings, app names, and dates are hidden on purpose:
label what the text says, not what you guess the rating was.

Pick one value in the `label` column for every row:
`negative`, `neutral`, `positive`, or `invalid`. Use `notes` for anything unusual
(e.g. "sarcasm?", "mostly English", "unsure between neutral/negative").

## Labels

### negative
A complaint, problem, frustration, or accusation, **even if it is polite**.
Includes bugs, failed transactions, lost balance, slow service, scam/fraud
accusations, and disappointment.

- **Requests framed as complaints are negative.** "kenapa limit saya turun",
  "tolong kembalikan saldo saya" describe a problem the user is unhappy about.

{table('negative')}

### neutral
A question, a feature request, or a factual statement **without clear emotion**.
Also **mixed reviews where neither side dominates**.

- A plain feature request ("tolong tambahkan fitur ...") is neutral; if it
  comes with frustration about a problem, it is negative.

{table('neutral')}

### positive
Praise, satisfaction, gratitude, or recommendation.

{table('positive')}

### invalid
Gibberish, unreadable text, or text unrelated to the app (e.g. a review meant
for a different product, random characters). Invalid rows are replaced from the
`reserve` sheet, so use this label sparingly.

{table('invalid')}

## Rules for hard cases

Examples in this section are illustrative, not taken from the data.

1. **Mixed reviews:** label the **dominant** sentiment. Use `neutral` only when
   praise and complaint are truly balanced.
   - "aplikasi bagus tapi sering error pas transfer, tolong diperbaiki" → the
     complaint dominates → `negative`.
2. **Polite tone does not change the label.** A courteous complaint is still `negative`.
3. **English or mixed language:** label normally.
4. **Emoji only or very short text:** label it if the meaning is clear
   ("mantap 👍" → `positive`); use `invalid` only if it is unreadable.
5. **Sarcasm:** label the intended meaning ("mantap, saldo hilang lagi" → `negative`)
   and add `sarcasm` in `notes`.
6. **Do not look up** the original review, rating, or app. Label the text as shown.

## Workflow

- Label the `gold` sheet top to bottom. Do not reorder or delete rows.
- When you mark a row `invalid`, leave it in place. Each invalid row is replaced
  by a reserve row from the same hidden stratum, so the import step (3b) will
  list exactly which `reserve` rows (R....) to label. You may also label the
  whole `reserve` sheet up front; unused reserve labels are ignored.
- Save the file as `.xlsx` with the same name.
"""


def main() -> None:
    """CLI entry point."""
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--config", type=Path, default=Path("configs/label.yaml"))
    parser.add_argument(
        "--guidelines-only",
        action="store_true",
        help="only re-render GUIDELINES.md from the existing gold key (no resampling, no overwrite)",
    )
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")

    cfg = load_config(args.config)
    logger.info("seed=%s config=%s", cfg["seed"], args.config)

    df = pd.read_parquet(cfg["input_path"])
    pool = build_pool(df, cfg["weak_label_map"])

    if args.guidelines_only:
        key = pd.read_parquet(cfg["gold"]["key_path"])
    else:
        pool.to_parquet(cfg["pool_output_path"], index=False)
        logger.info("pool: %d rows with weak_label -> %s", len(pool), cfg["pool_output_path"])
        key = sample_gold(pool, cfg)
        export_gold(key, cfg)

        splits = build_splits(pool, key, cfg)
        out_dir = Path(cfg["splits"]["dir"])
        out_dir.mkdir(parents=True, exist_ok=True)
        for name, frame in splits.items():
            frame.to_parquet(out_dir / f"{name}.parquet", index=False)
        logger.info("splits: %s -> %s", {k: len(v) for k, v in splits.items()}, out_dir)

        overlaps = assert_no_overlap(
            {
                "train": set(splits["train"]["reviewId"]),
                "val": set(splits["val"]["reviewId"]),
                "gold": set(key.loc[key["set"] == "gold", "reviewId"]),
                "reserve": set(key.loc[key["set"] == "reserve", "reviewId"]),
            }
        )
        print_report(pool, key, splits, overlaps)

    path = Path(cfg["guidelines"]["path"])
    path.write_text(render_guidelines(pool, key, cfg), encoding="utf-8")
    logger.info("wrote %s (%d examples)", path, len(cfg["guidelines"]["examples"]))


def print_report(
    pool: pd.DataFrame, key: pd.DataFrame, splits: dict[str, pd.DataFrame], overlaps: dict[str, int]
) -> None:
    """Print gold strata, split composition, and overlap checks."""
    print("\n== weak_label in pool ==")
    print(pool.groupby(["app", "weak_label"]).size().unstack().to_string())

    print("\n== gold / reserve strata (app x weak_label) ==")
    print(key.groupby(["set", "app", "weak_label"]).size().unstack().to_string())
    print("short reviews in gold:", int(key.loc[key["set"] == "gold", "is_short"].sum()),
          "| in reserve:", int(key.loc[key["set"] == "reserve", "is_short"].sum()))

    for name, frame in splits.items():
        print(f"\n== {name}: {len(frame)} rows ==")
        t = frame.groupby(["app", "weak_label"]).size().unstack()
        t["total"] = t.sum(axis=1)
        t.loc["total"] = t.sum()
        print(t.to_string())
    full = pd.concat(splits.values())
    print("\nlabel share (train+val):", full["weak_label"].value_counts(normalize=True).round(3).to_dict())
    weeks = full.groupby(["app", "week_start"]).size().unstack("app")
    weeks.index = weeks.index.strftime("%m-%d")
    print("\nrows per week (train+val):")
    print(weeks.to_string())
    print("\noverlap checks (all must be 0):", overlaps)


if __name__ == "__main__":
    main()
